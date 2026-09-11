from __future__ import annotations

import json
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

from zn_agent.core.browser_scene import BrowserFrameScene, BrowserScene, BrowserSceneTarget
from zn_agent.core.browser_scene_table import BrowserStructuredTable, BrowserTableCell, BrowserTableRow
from zn_agent.core.provider_bridge import build_resident_runtime

TASK = "把这个网站里的数据整理进我现在这个表里。"


class _FakeBrowser:
    def __init__(self, *, headers=None, rows=None, table_count=1):
        self.headers = list(headers or ["订单号", "客户", "金额"])
        self.rows = [list(row) for row in (rows or [["1002", "B", "99.50"], ["1003", "C", "8.00"]])]
        self.table_count = table_count
        self.observe_count = 0
        self.table_observe_count = 0
        self.url = "https://example.test/orders"

    def _session(self, session_id):
        if session_id != "session-1":
            raise RuntimeError("unknown managed browser session")
        return SimpleNamespace(identity=SimpleNamespace(provider="fake-managed-browser"))

    def observe_scene(self, session_id, *, page_id=""):
        if session_id != "session-1" or page_id != "page-1":
            raise RuntimeError("unknown managed browser page")
        self.observe_count += 1
        frame = BrowserFrameScene(
            frame_id="frame-main",
            parent_frame_id="",
            url=self.url,
            name="",
            is_main=True,
            observable=True,
        )
        targets = tuple(
            BrowserSceneTarget(
                target_id=f"table-{self.observe_count}-{index}",
                frame_id="frame-main",
                role="table",
                accessible_name=f"Orders {index}",
                visible=True,
                enabled=True,
                editable=False,
                tag="table",
            )
            for index in range(self.table_count)
        )
        return BrowserScene(
            scene_id=f"scene-{self.observe_count}",
            session_id=session_id,
            page_id=page_id,
            captured_at=f"capture-{self.observe_count}",
            url=self.url,
            title="Orders",
            frames=(frame,),
            targets=targets,
            truncated=False,
        )

    def observe_scene_table(self, session_id, target_id, *, page_id="", **kwargs):
        self.table_observe_count += 1
        table_rows = [
            BrowserTableRow(
                cells=tuple(BrowserTableCell(kind="header", text=value) for value in self.headers)
            )
        ]
        table_rows.extend(
            BrowserTableRow(
                cells=tuple(BrowserTableCell(kind="cell", text=value) for value in row)
            )
            for row in self.rows
        )
        return BrowserStructuredTable(
            session_id=session_id,
            page_id=page_id,
            target_id=target_id,
            captured_at=f"table-capture-{self.table_observe_count}",
            rows=tuple(table_rows),
            row_count_observed=len(table_rows),
            truncated=False,
        )


class BrowserSpreadsheetBehaviorTests(unittest.TestCase):
    @staticmethod
    def _workbook(path: Path, *, headers=None):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "销售"
        sheet.append(list(headers or ["订单号", "客户", "金额"]))
        sheet.append([1001, "A", 1200])
        workbook.save(path)

    def _resident(self, root: Path, thread_id: str, *, browser=None):
        db_path = root / f"zn-browser-sheet-{thread_id}-{os.getpid()}.db"
        resident = build_resident_runtime(config={"model": {}}, store_path=db_path)
        resident.managed_browser = browser or _FakeBrowser()
        self.addCleanup(resident.store.close)
        ledger = resident.work_ledger
        workspace = root / thread_id
        workspace.mkdir()
        ledger.create_thread(thread_id=thread_id)
        ledger.attach_workspace(thread_id, workspace)
        return resident, ledger, workspace

    @staticmethod
    def _payload(source: Path):
        return {
            "spreadsheet_path": str(source),
            "browser_session_id": "session-1",
            "page_id": "page-1",
        }

    @staticmethod
    def _run_to_terminal(resident, event_id: str):
        for _ in range(192):
            completed = resident.result_for(event_id)
            if completed is not None:
                return completed
            result = resident.live_once()
            if result is not None and result.event.event_id == event_id:
                return result
        raise AssertionError(f"event did not terminate: {resident.store.get_working_state()}")

    @staticmethod
    def _run_to_stage(resident, event_id: str, stage: str):
        for _ in range(128):
            state = resident.store.get_working_state()
            if state.current_event_id == event_id and state.stage == stage:
                return state
            result = resident.live_once()
            if result is not None:
                raise AssertionError(f"event terminated before {stage}: {result}")
        raise AssertionError(f"event did not reach {stage}")

    @staticmethod
    def _actions(resident, event_id: str):
        return [item for item in reversed(resident.body.recent_actions(512)) if item.event_id == event_id]

    def test_plain_user_intent_runs_zero_model_and_requires_fresh_browser_and_xlsx_evidence(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            browser = _FakeBrowser()
            resident, ledger, workspace = self._resident(root, "happy", browser=browser)
            source = workspace / "sales.xlsx"
            self._workbook(source)
            source_bytes = source.read_bytes()
            _, run = ledger.submit("happy", TASK, payload=self._payload(source))
            self.assertTrue(run.success, run)
            self.assertEqual(run.model_invocations, 0)
            self.assertEqual(source.read_bytes(), source_bytes)
            destination = workspace / "sales-webdata.xlsx"
            reopened = load_workbook(destination)
            self.assertEqual(
                [tuple(cell.value for cell in row) for row in reopened.active.iter_rows()],
                [
                    ("订单号", "客户", "金额"),
                    (1001, "A", 1200),
                    ("1002", "B", "99.50"),
                    ("1003", "C", "8.00"),
                ],
            )
            self.assertGreaterEqual(browser.observe_count, 3)
            actions = self._actions(resident, run.event.event_id)
            self.assertEqual(sum(item.kind == "append_xlsx_rows_copy" for item in actions), 1)
            self.assertGreaterEqual(sum(item.kind == "inspect_xlsx_append_target" for item in actions), 3)
            item = next(
                value
                for value in ledger.list_work_items("happy", limit=64)
                if "browser_spreadsheet_import:v1" in value.acceptance_criteria
            )
            evidence = json.loads(item.result)
            self.assertEqual(evidence["status"], "complete")
            self.assertEqual(evidence["ordered_headers"], ["订单号", "客户", "金额"])
            self.assertEqual(evidence["imported_row_count"], 2)
            self.assertTrue(evidence["fresh_browser_verification_fingerprint"])
            self.assertTrue(evidence["fresh_source_xlsx_identity_result"])
            self.assertTrue(evidence["fresh_destination_reopen_result"]["ready"])

    def test_missing_browser_or_unauthorized_spreadsheet_context_fails_closed(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            resident, ledger, workspace = self._resident(root, "missing")
            source = workspace / "sales.xlsx"
            self._workbook(source)
            _, run = ledger.submit("missing", TASK, payload={"spreadsheet_path": str(source)})
            self.assertFalse(run.success)
            self.assertIn("browser_context_missing", run.reason)
            self.assertFalse((workspace / "sales-webdata.xlsx").exists())

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            resident, ledger, workspace = self._resident(root, "unauthorized")
            outside = root / "outside.xlsx"
            self._workbook(outside)
            _, run = ledger.submit("unauthorized", TASK, payload=self._payload(outside))
            self.assertFalse(run.success)
            self.assertIn("spreadsheet_authority_missing", run.reason)
            self.assertFalse((workspace / "outside-webdata.xlsx").exists())

    def test_multiple_tables_and_header_mismatch_do_not_guess_or_write(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            browser = _FakeBrowser(table_count=2)
            resident, ledger, workspace = self._resident(root, "ambiguous", browser=browser)
            source = workspace / "sales.xlsx"
            self._workbook(source)
            _, run = ledger.submit("ambiguous", TASK, payload=self._payload(source))
            self.assertFalse(run.success)
            self.assertIn("ambiguous_browser_table", run.reason)
            self.assertFalse((workspace / "sales-webdata.xlsx").exists())

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            resident, ledger, workspace = self._resident(root, "headers")
            source = workspace / "sales.xlsx"
            self._workbook(source, headers=["客户", "订单号", "金额"])
            _, run = ledger.submit("headers", TASK, payload=self._payload(source))
            self.assertFalse(run.success)
            self.assertIn("spreadsheet_header_mismatch", run.reason)
            self.assertFalse((workspace / "sales-webdata.xlsx").exists())
            self.assertFalse(any(item.kind == "append_xlsx_rows_copy" for item in self._actions(resident, run.event.event_id)))

    def test_browser_and_source_drift_stop_before_destination_publish(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            browser = _FakeBrowser()
            resident, ledger, workspace = self._resident(root, "browser-drift", browser=browser)
            source = workspace / "sales.xlsx"
            self._workbook(source)
            _, event = ledger.start("browser-drift", TASK, payload=self._payload(source))
            self._run_to_stage(resident, event.event_id, "browser_spreadsheet_mutate")
            browser.rows[0][2] = "100.00"
            run = self._run_to_terminal(resident, event.event_id)
            self.assertFalse(run.success)
            self.assertIn("stale_browser_table_evidence", run.reason)
            self.assertFalse((workspace / "sales-webdata.xlsx").exists())
            self.assertFalse(any(item.kind == "append_xlsx_rows_copy" for item in self._actions(resident, event.event_id)))

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            resident, ledger, workspace = self._resident(root, "source-drift")
            source = workspace / "sales.xlsx"
            self._workbook(source)
            _, event = ledger.start("source-drift", TASK, payload=self._payload(source))
            self._run_to_stage(resident, event.event_id, "browser_spreadsheet_mutate")
            workbook = load_workbook(source)
            workbook.active.append([1009, "drift", 1])
            workbook.save(source)
            run = self._run_to_terminal(resident, event.event_id)
            self.assertFalse(run.success)
            self.assertIn("stale_source_evidence", run.reason)
            self.assertFalse((workspace / "sales-webdata.xlsx").exists())

    def test_output_collision_and_complex_workbook_fail_closed(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            resident, ledger, workspace = self._resident(root, "collision")
            source = workspace / "sales.xlsx"
            self._workbook(source)
            destination = workspace / "sales-webdata.xlsx"
            destination.write_bytes(b"sentinel")
            _, run = ledger.submit("collision", TASK, payload=self._payload(source))
            self.assertFalse(run.success)
            self.assertIn("output_collision", run.reason)
            self.assertEqual(destination.read_bytes(), b"sentinel")

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            resident, ledger, workspace = self._resident(root, "formula")
            source = workspace / "formula.xlsx"
            self._workbook(source)
            workbook = load_workbook(source)
            workbook.active["C2"] = "=1+1"
            workbook.save(source)
            _, run = ledger.submit("formula", TASK, payload=self._payload(source))
            self.assertFalse(run.success)
            self.assertIn("unsupported_workbook_structure", run.reason)
            self.assertFalse((workspace / "formula-webdata.xlsx").exists())


if __name__ == "__main__":
    unittest.main(verbosity=2)

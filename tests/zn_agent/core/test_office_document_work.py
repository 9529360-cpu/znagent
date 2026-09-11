from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Pt

from zn_agent.core.office_document import (
    inspect_docx,
    inspect_docx_completion_targets,
    replace_visible_span_across_runs,
    write_docx_completion_copy,
    write_docx_copy,
)


class OfficeDocumentWorkTests(unittest.TestCase):
    @staticmethod
    def _contract(path: Path, *, table_target: bool = False, duplicate: bool = False) -> None:
        doc = Document()
        doc.add_paragraph("合同编号：A-100")
        if table_target:
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "条款"
            table.cell(0, 1).text = "内容"
            paragraph = table.cell(1, 1).paragraphs[0]
        else:
            paragraph = doc.add_paragraph()
        paragraph.add_run("付款日期：")
        year = paragraph.add_run("2026年")
        year.bold = True
        year.font.size = Pt(11)
        month_day = paragraph.add_run("9月15日")
        month_day.italic = True
        if duplicate:
            doc.add_paragraph("付款日期：2026年9月15日")
        doc.add_paragraph("其他条款保持不变。")
        doc.save(path)

    @staticmethod
    def _completion_document(path: Path, *, targets: int = 3) -> None:
        doc = Document()
        doc.add_paragraph("方案名称：Project Northstar")
        paragraph = doc.add_paragraph("费用：")
        first = paragraph.add_run("【待")
        first.bold = True
        first.font.size = Pt(11)
        second = paragraph.add_run("补充：官方价格是多少？】")
        second.italic = True
        if targets >= 2:
            paragraph.add_run("；发布日期：")
            date_run = paragraph.add_run("【待补充：正式发布日期是什么？】")
            date_run.font.size = Pt(12)
        if targets >= 3:
            table = doc.add_table(rows=1, cols=1)
            cell = table.cell(0, 0).paragraphs[0]
            cell.add_run("可用地区：")
            region = cell.add_run("【待补充：首发支持哪些国家？】")
            region.bold = True
        if targets >= 4:
            doc.add_paragraph("额外：【待补充：第四个问题是什么？】")
        doc.add_paragraph("审批流程：保持原样。")
        doc.save(path)

    def test_run_local_and_cross_run_replacement_preserve_run_shape(self) -> None:
        doc = Document()
        paragraph = doc.add_paragraph()
        run = paragraph.add_run("付款日期：2026年9月15日；其他")
        run.bold = True
        start = paragraph.text.index("2026")
        replace_visible_span_across_runs(paragraph, start, start + len("2026年9月15日"), "2026年9月20日")
        self.assertEqual(paragraph.text, "付款日期：2026年9月20日；其他")
        self.assertTrue(paragraph.runs[0].bold)

        paragraph = doc.add_paragraph()
        paragraph.add_run("付款日期：")
        year = paragraph.add_run("2026年")
        year.bold = True
        month_day = paragraph.add_run("9月15日")
        month_day.italic = True
        start = paragraph.text.index("2026")
        replace_visible_span_across_runs(paragraph, start, len(paragraph.text), "2026年9月20日")
        self.assertEqual(paragraph.text, "付款日期：2026年9月20日")
        self.assertTrue(paragraph.runs[1].bold)
        self.assertEqual(paragraph.runs[2].text, "")
        self.assertTrue(paragraph.runs[2].italic)

    def test_real_docx_copy_reopens_preserves_source_and_table_cell_target(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "contract.docx"
            destination = root / "project.docx"
            self._contract(source, table_target=True)
            before = source.read_bytes()
            inspected = inspect_docx(source)
            self.assertTrue(inspected["ready"], inspected)
            self.assertEqual(inspected["payment_date_occurrence_count"], 1)
            self.assertTrue(inspected["payment_date_occurrences"][0]["location"].startswith("table:"))
            result = write_docx_copy(
                source,
                destination,
                replacement_date="2026年9月20日",
                precondition_identity=inspected["identity"],
            )
            self.assertTrue(result["source_unchanged"])
            self.assertTrue(result["destination_reopened"])
            self.assertTrue(result["format_preserved"])
            self.assertEqual(source.read_bytes(), before)
            reopened = Document(destination)
            texts = [
                paragraph.text
                for table in reopened.tables
                for row in table.rows
                for cell in row.cells
                for paragraph in cell.paragraphs
            ]
            self.assertIn("付款日期：2026年9月20日", texts)

    def test_zero_multiple_drift_collision_and_unsupported_structure_fail_closed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            no_target = root / "none.docx"
            Document().save(no_target)
            self.assertEqual(inspect_docx(no_target)["payment_date_occurrence_count"], 0)

            repeated = root / "repeat.docx"
            self._contract(repeated, duplicate=True)
            repeated_info = inspect_docx(repeated)
            self.assertEqual(repeated_info["payment_date_occurrence_count"], 2)
            with self.assertRaises(ValueError):
                write_docx_copy(
                    repeated,
                    root / "repeat-out.docx",
                    replacement_date="2026年9月20日",
                    precondition_identity=repeated_info["identity"],
                )

            drift = root / "drift.docx"
            self._contract(drift)
            drift_info = inspect_docx(drift)
            doc = Document(drift)
            doc.add_paragraph("外部更新")
            doc.save(drift)
            with self.assertRaisesRegex(RuntimeError, "stale_source_evidence"):
                write_docx_copy(
                    drift,
                    root / "drift-out.docx",
                    replacement_date="2026年9月20日",
                    precondition_identity=drift_info["identity"],
                )
            self.assertFalse((root / "drift-out.docx").exists())

            collision = root / "collision.docx"
            self._contract(collision)
            collision_info = inspect_docx(collision)
            existing = root / "exists.docx"
            existing.write_bytes(b"sentinel")
            with self.assertRaises(FileExistsError):
                write_docx_copy(
                    collision,
                    existing,
                    replacement_date="2026年9月20日",
                    precondition_identity=collision_info["identity"],
                )
            self.assertEqual(existing.read_bytes(), b"sentinel")

            merged = root / "merged.docx"
            doc = Document()
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).merge(table.cell(0, 1)).text = "付款日期：2026年9月15日"
            doc.save(merged)
            blocked = inspect_docx(merged)
            self.assertFalse(blocked["ready"])
            self.assertEqual(blocked["blocker"], "unsupported_document_structure")

    def test_completion_inspection_binds_split_run_body_and_table_targets(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "proposal.docx"
            self._completion_document(source, targets=3)
            first = inspect_docx_completion_targets(source)
            second = inspect_docx_completion_targets(source)
            self.assertTrue(first["ready"], first)
            self.assertEqual(first["target_count"], 3)
            self.assertEqual(
                [item["target_id"] for item in first["targets"]],
                [item["target_id"] for item in second["targets"]],
            )
            self.assertTrue(all(item["run_mappable"] for item in first["targets"]))
            self.assertTrue(first["targets"][0]["location"].startswith("body/"))
            self.assertTrue(first["targets"][2]["location"].startswith("table:"))
            self.assertLessEqual(max(len(item["question"]) for item in first["targets"]), 300)
            self.assertLessEqual(max(len(item["context_text"]) for item in first["targets"]), 800)

    def test_completion_copy_replaces_all_targets_preserving_source_unrelated_text_and_run_format(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "proposal.docx"
            destination = root / "proposal-completed.docx"
            self._completion_document(source, targets=3)
            source_bytes = source.read_bytes()
            before_doc = Document(source)
            body_formats = [
                (run.bold, run.italic, run.font.size.pt if run.font.size else None)
                for run in before_doc.paragraphs[1].runs
            ]
            table_formats = [
                (run.bold, run.italic, run.font.size.pt if run.font.size else None)
                for run in before_doc.tables[0].cell(0, 0).paragraphs[0].runs
            ]
            inspected = inspect_docx_completion_targets(source)
            ids = [item["target_id"] for item in inspected["targets"]]
            result = write_docx_completion_copy(
                source,
                destination,
                replacements=[
                    {"target_id": ids[0], "text": "官方价格为 99 美元。"},
                    {"target_id": ids[1], "text": "正式发布日期为 2026年9月1日。"},
                    {"target_id": ids[2], "text": "首发支持马来西亚和新加坡。"},
                ],
                precondition_identity=inspected["identity"],
            )
            self.assertEqual(source.read_bytes(), source_bytes)
            self.assertTrue(result["source_unchanged"])
            self.assertTrue(result["destination_reopened"])
            self.assertTrue(result["non_target_text_preserved"])
            self.assertTrue(result["run_topology_preserved"])
            self.assertTrue(result["format_preserved"])
            self.assertTrue(result["all_placeholders_replaced"])
            self.assertEqual(result["target_count_before"], 3)
            self.assertEqual(result["target_count_after"], 0)

            reopened_info = inspect_docx_completion_targets(destination)
            self.assertTrue(reopened_info["ready"], reopened_info)
            self.assertEqual(reopened_info["target_count"], 0)
            reopened = Document(destination)
            self.assertEqual(
                reopened.paragraphs[1].text,
                "费用：官方价格为 99 美元。；发布日期：正式发布日期为 2026年9月1日。",
            )
            self.assertEqual(reopened.tables[0].cell(0, 0).paragraphs[0].text, "可用地区：首发支持马来西亚和新加坡。")
            self.assertEqual(reopened.paragraphs[-1].text, "审批流程：保持原样。")
            self.assertEqual(
                [
                    (run.bold, run.italic, run.font.size.pt if run.font.size else None)
                    for run in reopened.paragraphs[1].runs
                ],
                body_formats,
            )
            self.assertEqual(
                [
                    (run.bold, run.italic, run.font.size.pt if run.font.size else None)
                    for run in reopened.tables[0].cell(0, 0).paragraphs[0].runs
                ],
                table_formats,
            )

    def test_completion_boundaries_target_drift_source_drift_collision_and_unsupported_fail_closed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            overflow = root / "overflow.docx"
            self._completion_document(overflow, targets=4)
            blocked = inspect_docx_completion_targets(overflow)
            self.assertFalse(blocked["ready"])
            self.assertEqual(blocked["blocker"], "completion_target_count_out_of_bounds")

            malformed = root / "malformed.docx"
            doc = Document()
            doc.add_paragraph("【待补充:缺少全角冒号】")
            doc.save(malformed)
            malformed_info = inspect_docx_completion_targets(malformed)
            self.assertFalse(malformed_info["ready"])
            self.assertEqual(malformed_info["blocker"], "invalid_completion_placeholder")

            source = root / "target-drift.docx"
            self._completion_document(source, targets=1)
            bound = inspect_docx_completion_targets(source)
            with self.assertRaisesRegex(RuntimeError, "target_drift"):
                write_docx_completion_copy(
                    source,
                    root / "target-drift-out.docx",
                    replacements=[{"target_id": "doc-target-not-bound", "text": "x"}],
                    precondition_identity=bound["identity"],
                )
            self.assertFalse((root / "target-drift-out.docx").exists())

            drift = root / "source-drift.docx"
            self._completion_document(drift, targets=1)
            drift_bound = inspect_docx_completion_targets(drift)
            doc = Document(drift)
            doc.add_paragraph("外部更新")
            doc.save(drift)
            with self.assertRaisesRegex(RuntimeError, "stale_source_evidence"):
                write_docx_completion_copy(
                    drift,
                    root / "source-drift-out.docx",
                    replacements=[{"target_id": drift_bound["targets"][0]["target_id"], "text": "x"}],
                    precondition_identity=drift_bound["identity"],
                )
            self.assertFalse((root / "source-drift-out.docx").exists())

            collision = root / "collision-source.docx"
            self._completion_document(collision, targets=1)
            collision_bound = inspect_docx_completion_targets(collision)
            existing = root / "collision-source-completed.docx"
            existing.write_bytes(b"sentinel")
            with self.assertRaises(FileExistsError):
                write_docx_completion_copy(
                    collision,
                    existing,
                    replacements=[{"target_id": collision_bound["targets"][0]["target_id"], "text": "x"}],
                    precondition_identity=collision_bound["identity"],
                )
            self.assertEqual(existing.read_bytes(), b"sentinel")

            merged = root / "merged-completion.docx"
            doc = Document()
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).merge(table.cell(0, 1)).text = "【待补充：价格是什么？】"
            doc.save(merged)
            merged_info = inspect_docx_completion_targets(merged)
            self.assertFalse(merged_info["ready"])
            self.assertEqual(merged_info["blocker"], "unsupported_document_structure")


if __name__ == "__main__":
    unittest.main()

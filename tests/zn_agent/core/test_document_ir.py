from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from zn_agent.core.document_ir import (
    Element,
    flatten_elements,
    parse_document,
    render_table,
)


def _make_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("Report Title", level=1)
    document.add_paragraph("First paragraph with some text.")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Age"
    table.cell(1, 0).text = "Alice"
    table.cell(1, 1).text = "30"
    table.cell(2, 0).text = "Bob"
    table.cell(2, 1).text = "25"
    document.add_paragraph("After the table.")
    document.save(path)


def _make_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "People"
    sheet.append(["Name", "Age"])
    sheet.append(["Alice", 30])
    sheet.append(["Bob", 25])
    workbook.save(path)


class DocumentIRTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(self.enterContext(tempfile.TemporaryDirectory()))

    def test_docx_parses_in_stream_order(self):
        path = self.tmp / "doc.docx"
        _make_docx(path)
        elements = parse_document(path)
        kinds = [el.type for el in elements]
        self.assertEqual(kinds, ["heading", "paragraph", "table", "paragraph"])
        self.assertEqual(elements[0].text, "Report Title")
        self.assertEqual(elements[0].meta["level"], 1)

    def test_docx_table_keeps_rows(self):
        path = self.tmp / "doc.docx"
        _make_docx(path)
        table = next(el for el in parse_document(path) if el.type == "table")
        rows = table.meta["rows"]
        self.assertEqual(rows[0], ["Name", "Age"])
        self.assertEqual(rows[1], ["Alice", "30"])

    def test_xlsx_sheet_becomes_table(self):
        path = self.tmp / "sheet.xlsx"
        _make_xlsx(path)
        elements = parse_document(path)
        self.assertEqual(len(elements), 1)
        self.assertEqual(elements[0].type, "table")
        self.assertEqual(elements[0].meta["sheet"], "People")
        self.assertEqual(elements[0].meta["rows"][0], ["Name", "Age"])

    def test_text_headings_and_lists(self):
        path = self.tmp / "notes.md"
        path.write_text("# Title\n\n- item one\n- item two\n\nBody text.\n", encoding="utf-8")
        elements = parse_document(path)
        kinds = [el.type for el in elements]
        self.assertEqual(kinds, ["heading", "list_item", "list_item", "paragraph"])

    def test_unsupported_format_is_not_silent(self):
        path = self.tmp / "weird.xyz"
        path.write_bytes(b"nope")
        elements = parse_document(path)
        self.assertEqual(len(elements), 1)
        self.assertIn("unsupported", elements[0].text)

    def test_render_table_markdown(self):
        element = Element(type="table", meta={"rows": [["A", "B"], ["1", "2"]]})
        rendered = render_table(element)
        self.assertIn("A | B", rendered)
        self.assertIn("1 | 2", rendered)

    def test_flatten_preserves_order(self):
        elements = [
            Element(type="paragraph", text="hello"),
            Element(type="table", meta={"rows": [["H1", "H2"], ["v1", "v2"]]}),
            Element(type="paragraph", text="world"),
        ]
        flat = flatten_elements(elements)
        self.assertLess(flat.index("hello"), flat.index("H1 | H2"))
        self.assertLess(flat.index("H1 | H2"), flat.index("world"))


if __name__ == "__main__":
    unittest.main()

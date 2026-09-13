from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Pt

from zn_agent.core.office_document import inspect_docx, write_docx_copy


class OfficeDocumentHeaderFooterWorkTests(unittest.TestCase):
    @staticmethod
    def _add_payment_date(paragraph) -> None:
        paragraph.add_run("付款日期：")
        year = paragraph.add_run("2026年")
        year.bold = True
        year.font.size = Pt(11)
        month_day = paragraph.add_run("9月15日")
        month_day.italic = True

    @classmethod
    def _header_contract(cls, path: Path) -> None:
        doc = Document()
        doc.add_paragraph("正文条款保持不变。")
        header = doc.sections[0].header
        header.is_linked_to_previous = False
        cls._add_payment_date(header.paragraphs[0])
        doc.save(path)

    @classmethod
    def _footer_contract(cls, path: Path) -> None:
        doc = Document()
        doc.add_paragraph("正文条款保持不变。")
        footer = doc.sections[0].footer
        footer.is_linked_to_previous = False
        cls._add_payment_date(footer.paragraphs[0])
        doc.save(path)

    def test_header_payment_date_copy_preserves_source_body_and_run_format(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "header-contract.docx"
            destination = root / "header-contract-updated.docx"
            self._header_contract(source)
            source_bytes = source.read_bytes()

            inspected = inspect_docx(source)
            self.assertTrue(inspected["ready"], inspected)
            self.assertEqual(inspected["payment_date_occurrence_count"], 1)
            self.assertTrue(
                inspected["payment_date_occurrences"][0]["location"].startswith(
                    "header:section:0/default/"
                )
            )

            result = write_docx_copy(
                source,
                destination,
                replacement_date="2026年9月20日",
                precondition_identity=inspected["identity"],
            )

            self.assertEqual(source.read_bytes(), source_bytes)
            self.assertTrue(result["source_unchanged"])
            self.assertTrue(result["destination_reopened"])
            self.assertTrue(result["format_preserved"])
            self.assertTrue(str(result["target_location"]).startswith("header:"))

            reopened = Document(destination)
            paragraph = reopened.sections[0].header.paragraphs[0]
            self.assertEqual(paragraph.text, "付款日期：2026年9月20日")
            self.assertTrue(paragraph.runs[1].bold)
            self.assertAlmostEqual(paragraph.runs[1].font.size.pt, 11.0)
            self.assertTrue(paragraph.runs[2].italic)
            self.assertEqual(paragraph.runs[2].text, "")
            self.assertEqual(reopened.paragraphs[0].text, "正文条款保持不变。")

    def test_footer_payment_date_copy_is_bounded_and_reopens(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "footer-contract.docx"
            destination = root / "footer-contract-updated.docx"
            self._footer_contract(source)

            inspected = inspect_docx(source)
            self.assertTrue(inspected["ready"], inspected)
            self.assertEqual(inspected["payment_date_occurrence_count"], 1)
            self.assertTrue(
                inspected["payment_date_occurrences"][0]["location"].startswith(
                    "footer:section:0/default/"
                )
            )

            result = write_docx_copy(
                source,
                destination,
                replacement_date="2026年9月20日",
                precondition_identity=inspected["identity"],
            )
            self.assertTrue(result["destination_reopened"])
            self.assertTrue(str(result["target_location"]).startswith("footer:"))
            reopened = Document(destination)
            self.assertEqual(
                reopened.sections[0].footer.paragraphs[0].text,
                "付款日期：2026年9月20日",
            )

    def test_inactive_first_page_header_target_fails_closed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "inactive-first-header.docx"
            doc = Document()
            doc.add_paragraph("正文")
            section = doc.sections[0]
            self.assertFalse(section.different_first_page_header_footer)
            first = section.first_page_header
            first.is_linked_to_previous = False
            self._add_payment_date(first.paragraphs[0])
            doc.save(source)

            inspected = inspect_docx(source)
            self.assertFalse(inspected["ready"])
            self.assertEqual(inspected["blocker"], "unsupported_document_structure")
            self.assertIn("inactive", inspected["detail"])

    def test_body_and_header_targets_remain_ambiguous_for_mutation(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "duplicate-story-contract.docx"
            destination = root / "duplicate-story-contract-updated.docx"
            doc = Document()
            body = doc.add_paragraph()
            self._add_payment_date(body)
            header = doc.sections[0].header
            header.is_linked_to_previous = False
            self._add_payment_date(header.paragraphs[0])
            doc.save(source)

            inspected = inspect_docx(source)
            self.assertTrue(inspected["ready"], inspected)
            self.assertEqual(inspected["payment_date_occurrence_count"], 2)
            with self.assertRaises(ValueError):
                write_docx_copy(
                    source,
                    destination,
                    replacement_date="2026年9月20日",
                    precondition_identity=inspected["identity"],
                )
            self.assertFalse(destination.exists())


if __name__ == "__main__":
    unittest.main()

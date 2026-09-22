from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from zn_agent.core.presentation_document import (
    MAX_BULLETS_PER_SLIDE,
    MAX_TITLE_CHARS,
    create_pptx_from_outline,
    inspect_pptx,
)


class PresentationDocumentTests(unittest.TestCase):
    def test_create_outline_reopens_and_verifies_exact_visible_structure(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            destination = Path(tmp) / "review.pptx"
            result = create_pptx_from_outline(
                destination,
                title="Q3 Review",
                subtitle="ZN generated",
                slides=[
                    {"title": "Agenda", "bullets": ["Revenue", "Risks"]},
                    {"title": "Next steps", "bullets": ["Ship", "Verify"]},
                ],
            )
            self.assertTrue(result["verified"])
            self.assertEqual(result["slide_count"], 3)
            self.assertTrue(result["structure_fingerprint"])
            inspection = inspect_pptx(destination)
            self.assertTrue(inspection["ready"], inspection)
            self.assertEqual(
                [slide["texts"] for slide in inspection["slides"]],
                [
                    ["Q3 Review", "ZN generated"],
                    ["Agenda", "Revenue", "Risks"],
                    ["Next steps", "Ship", "Verify"],
                ],
            )
            reopened = Presentation(destination)
            self.assertEqual(len(reopened.slides), 3)
            self.assertEqual(reopened.slides[1].shapes.title.text, "Agenda")

    def test_inspection_extracts_table_text_without_powerpoint(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "table.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.shapes.title.text = "Metrics"
            table = slide.shapes.add_table(
                2, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)
            ).table
            table.cell(0, 0).text = "Name"
            table.cell(0, 1).text = "Value"
            table.cell(1, 0).text = "Users"
            table.cell(1, 1).text = "42"
            presentation.save(path)

            inspection = inspect_pptx(path)
            self.assertTrue(inspection["ready"], inspection)
            self.assertEqual(inspection["slide_count"], 1)
            self.assertEqual(
                inspection["slides"][0]["texts"],
                ["Metrics", "Name", "Value", "Users", "42"],
            )

    def test_collision_and_outline_bounds_fail_closed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            collision = root / "exists.pptx"
            collision.write_bytes(b"sentinel")
            with self.assertRaisesRegex(FileExistsError, "output_collision"):
                create_pptx_from_outline(collision, title="Title")
            self.assertEqual(collision.read_bytes(), b"sentinel")

            with self.assertRaisesRegex(ValueError, "presentation title exceeds"):
                create_pptx_from_outline(
                    root / "too-long.pptx",
                    title="x" * (MAX_TITLE_CHARS + 1),
                )
            with self.assertRaisesRegex(ValueError, "supports at most"):
                create_pptx_from_outline(
                    root / "too-many-bullets.pptx",
                    title="Title",
                    slides=[
                        {
                            "title": "Slide",
                            "bullets": ["item"] * (MAX_BULLETS_PER_SLIDE + 1),
                        }
                    ],
                )

    def test_invalid_package_and_wrong_extension_are_blocked(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            malformed = root / "bad.pptx"
            malformed.write_bytes(b"not a zip package")
            inspection = inspect_pptx(malformed)
            self.assertFalse(inspection["ready"])
            self.assertEqual(inspection["blocker"], "unsupported_presentation_structure")

            wrong = root / "deck.pptm"
            wrong.write_bytes(b"anything")
            inspection = inspect_pptx(wrong)
            self.assertFalse(inspection["ready"])
            self.assertIn("only standard .pptx", inspection["detail"])

            unsafe = root / "unsafe.pptx"
            presentation = Presentation()
            presentation.slides.add_slide(presentation.slide_layouts[0])
            presentation.save(unsafe)
            with zipfile.ZipFile(unsafe, mode="a") as package:
                package.writestr("ppt/vbaProject.bin", b"synthetic macro payload")
            inspection = inspect_pptx(unsafe)
            self.assertFalse(inspection["ready"])
            self.assertIn("unsupported PPTX package part", inspection["detail"])


if __name__ == "__main__":
    unittest.main()

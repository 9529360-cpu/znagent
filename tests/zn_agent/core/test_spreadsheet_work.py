from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference

from zn_agent.core.spreadsheet_work import (
    AMOUNT_NUMBER_FORMAT,
    append_xlsx_rows_copy,
    inspect_xlsx,
    inspect_xlsx_append_target,
    write_xlsx_copy,
)


class SpreadsheetWorkTests(unittest.TestCase):
    @staticmethod
    def _sales(path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "销售"
        sheet.append(["订单号", "客户", "金额"])
        rows = [(1001, "A", 1200), (1002, "B", 99.5), (1001, "A", 1200), (1003, "C", 8), (1002, "B", 99.5)]
        formats = ["0", "0.0", "0.00", "General", "#,##0"]
        for index, row in enumerate(rows, start=2):
            sheet.append(row)
            sheet.cell(index, 3).number_format = formats[index - 2]
        workbook.save(path)

    @staticmethod
    def _append_target(path: Path, *, with_data: bool = True) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "销售"
        sheet.append(["订单号", "客户", "金额"])
        if with_data:
            sheet.append([1001, "A", 1200])
            sheet.cell(2, 1).number_format = "0"
            sheet.cell(2, 3).number_format = "#,##0.00"
        workbook.save(path)

    def test_exact_dedupe_stable_retention_numeric_semantics_and_reopen(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "sales.xlsx"
            destination = root / "sales-cleaned.xlsx"
            self._sales(source)
            source_bytes = source.read_bytes()
            before = inspect_xlsx(source)
            self.assertTrue(before["ready"], before)
            self.assertEqual(before["duplicate_row_count"], 2)
            result = write_xlsx_copy(source, destination, precondition_identity=before["identity"])
            self.assertEqual(result["removed_duplicate_rows"], 2)
            self.assertTrue(result["stable_first_row_retention"])
            self.assertTrue(result["amount_values_semantically_unchanged"])
            self.assertTrue(result["amount_cells_numeric"])
            self.assertTrue(result["non_target_values_unchanged"])
            self.assertEqual(source.read_bytes(), source_bytes)
            reopened = load_workbook(destination)
            sheet = reopened.active
            self.assertEqual(
                [tuple(cell.value for cell in row) for row in sheet.iter_rows()],
                [("订单号", "客户", "金额"), (1001, "A", 1200), (1002, "B", 99.5), (1003, "C", 8)],
            )
            self.assertEqual(
                [sheet.cell(row, 3).number_format for row in range(2, sheet.max_row + 1)],
                [AMOUNT_NUMBER_FORMAT] * 3,
            )
            self.assertTrue(all(sheet.cell(row, 3).data_type == "n" for row in range(2, sheet.max_row + 1)))

    def test_ambiguous_amount_formula_complex_drift_and_collision_fail_closed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            ambiguous = root / "ambiguous.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["订单", "含税金额", "未税金额"])
            sheet.append([1, 10, 9])
            workbook.save(ambiguous)
            info = inspect_xlsx(ambiguous)
            self.assertFalse(info["ready"])
            self.assertEqual(info["blocker"], "ambiguous_amount_column")

            formula = root / "formula.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["订单", "金额"])
            sheet.append([1, "=1+1"])
            workbook.save(formula)
            blocked = inspect_xlsx(formula)
            self.assertFalse(blocked["ready"])
            self.assertEqual(blocked["blocker"], "unsupported_workbook_structure")

            charted = root / "chart.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["订单", "金额"])
            sheet.append([1, 10])
            chart = BarChart()
            chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2))
            sheet.add_chart(chart, "E2")
            workbook.save(charted)
            blocked = inspect_xlsx(charted)
            self.assertFalse(blocked["ready"])
            self.assertEqual(blocked["blocker"], "unsupported_workbook_structure")

            drift = root / "drift.xlsx"
            self._sales(drift)
            baseline = inspect_xlsx(drift)
            workbook = load_workbook(drift)
            workbook.active.append([1004, "D", 5])
            workbook.save(drift)
            with self.assertRaisesRegex(RuntimeError, "stale_source_evidence"):
                write_xlsx_copy(drift, root / "drift-cleaned.xlsx", precondition_identity=baseline["identity"])
            self.assertFalse((root / "drift-cleaned.xlsx").exists())

            source = root / "collision.xlsx"
            self._sales(source)
            baseline = inspect_xlsx(source)
            destination = root / "collision-cleaned.xlsx"
            destination.write_bytes(b"sentinel")
            with self.assertRaisesRegex(FileExistsError, "output_collision"):
                write_xlsx_copy(source, destination, precondition_identity=baseline["identity"])
            self.assertEqual(destination.read_bytes(), b"sentinel")

    def test_append_target_accepts_header_only_and_existing_rectangular_data(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            header_only = root / "header-only.xlsx"
            existing = root / "existing.xlsx"
            self._append_target(header_only, with_data=False)
            self._append_target(existing, with_data=True)
            first = inspect_xlsx_append_target(header_only)
            second = inspect_xlsx_append_target(existing)
            self.assertTrue(first["ready"], first)
            self.assertEqual(first["existing_row_count"], 0)
            self.assertTrue(second["ready"], second)
            self.assertEqual(second["existing_row_count"], 1)
            self.assertEqual(second["headers"], ["订单号", "客户", "金额"])
            self.assertTrue(second["semantic_fingerprint"])

    def test_append_copy_preserves_existing_cells_and_source_and_reopens_exact_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "sales.xlsx"
            destination = root / "sales-webdata.xlsx"
            self._append_target(source)
            source_bytes = source.read_bytes()
            before = inspect_xlsx_append_target(source)
            rows = [["1002", "B", "99.50"], ["=1+1", "C", "8.00"]]
            result = append_xlsx_rows_copy(
                source,
                destination,
                precondition_identity=before["identity"],
                headers=["订单号", "客户", "金额"],
                rows=rows,
            )
            self.assertEqual(source.read_bytes(), source_bytes)
            self.assertTrue(result["source_unchanged"])
            self.assertTrue(result["destination_reopened"])
            self.assertTrue(result["existing_content_unchanged"])
            self.assertTrue(result["appended_rows_exact"])
            self.assertEqual(result["existing_semantic_fingerprint_before"], result["existing_semantic_fingerprint_after"])
            reopened = load_workbook(destination)
            sheet = reopened.active
            self.assertEqual(sheet.title, "销售")
            self.assertEqual(
                [tuple(cell.value for cell in row) for row in sheet.iter_rows()],
                [
                    ("订单号", "客户", "金额"),
                    (1001, "A", 1200),
                    ("1002", "B", "99.50"),
                    ("=1+1", "C", "8.00"),
                ],
            )
            self.assertEqual(sheet.cell(2, 1).data_type, "n")
            self.assertEqual(sheet.cell(2, 1).number_format, "0")
            self.assertEqual(sheet.cell(2, 3).number_format, "#,##0.00")
            self.assertEqual(sheet.cell(3, 1).data_type, "s")
            self.assertEqual(sheet.cell(4, 1).data_type, "s")
            final = inspect_xlsx_append_target(destination)
            self.assertEqual(final["semantic_fingerprint"], result["destination_semantic_fingerprint"])
            self.assertEqual(final["existing_row_count"], 3)

    def test_append_header_malformed_row_source_drift_collision_and_complex_fail_closed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "sales.xlsx"
            self._append_target(source)
            before = inspect_xlsx_append_target(source)

            with self.assertRaisesRegex(ValueError, "spreadsheet_header_mismatch"):
                append_xlsx_rows_copy(
                    source,
                    root / "header-mismatch.xlsx",
                    precondition_identity=before["identity"],
                    headers=["客户", "订单号", "金额"],
                    rows=[["B", "1002", "99.50"]],
                )
            self.assertFalse((root / "header-mismatch.xlsx").exists())

            with self.assertRaisesRegex(ValueError, "rectangular"):
                append_xlsx_rows_copy(
                    source,
                    root / "malformed.xlsx",
                    precondition_identity=before["identity"],
                    headers=["订单号", "客户", "金额"],
                    rows=[["1002", "B"]],
                )
            self.assertFalse((root / "malformed.xlsx").exists())

            workbook = load_workbook(source)
            workbook.active.append([1009, "drift", 1])
            workbook.save(source)
            with self.assertRaisesRegex(RuntimeError, "stale_source_evidence"):
                append_xlsx_rows_copy(
                    source,
                    root / "drift-webdata.xlsx",
                    precondition_identity=before["identity"],
                    headers=["订单号", "客户", "金额"],
                    rows=[["1002", "B", "99.50"]],
                )
            self.assertFalse((root / "drift-webdata.xlsx").exists())

            current = inspect_xlsx_append_target(source)
            collision = root / "sales-webdata.xlsx"
            collision.write_bytes(b"sentinel")
            with self.assertRaisesRegex(FileExistsError, "output_collision"):
                append_xlsx_rows_copy(
                    source,
                    collision,
                    precondition_identity=current["identity"],
                    headers=["订单号", "客户", "金额"],
                    rows=[["1002", "B", "99.50"]],
                )
            self.assertEqual(collision.read_bytes(), b"sentinel")

            formula = root / "formula-append.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["订单号", "客户", "金额"])
            sheet.append([1, "A", "=1+1"])
            workbook.save(formula)
            info = inspect_xlsx_append_target(formula)
            self.assertFalse(info["ready"])
            self.assertEqual(info["blocker"], "unsupported_workbook_structure")

            charted = root / "chart-append.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["订单号", "客户", "金额"])
            sheet.append([1, "A", 10])
            chart = BarChart()
            chart.add_data(Reference(sheet, min_col=3, min_row=1, max_row=2))
            sheet.add_chart(chart, "E2")
            workbook.save(charted)
            info = inspect_xlsx_append_target(charted)
            self.assertFalse(info["ready"])
            self.assertEqual(info["blocker"], "unsupported_workbook_structure")


if __name__ == "__main__":
    unittest.main()

from __future__ import annotations

"""End-to-end payment-date round trip in DOCX header/footer stories.

Builds a real .docx on disk with payment dates in the header and the footer,
then drives the production path inspect_docx -> write_docx_copy -> reopen,
proving the date is replaced in the right story, the source file is untouched,
and the destination re-inspects cleanly. No mocks, no work-loop scaffolding.
"""

import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Pt

from zn_agent.core.office_document import inspect_docx, write_docx_copy


class HeaderFooterPaymentDateE2ETests(unittest.TestCase):
    @staticmethod
    def _add_payment_date(paragraph, date: str = "2026年9月15日") -> None:
        paragraph.add_run("付款日期：")
        year = paragraph.add_run(date[:5])
        year.bold = True
        year.font.size = Pt(11)
        month_day = paragraph.add_run(date[5:])
        month_day.italic = True

    @classmethod
    def _fixture(cls, path: Path) -> None:
        doc = Document()
        doc.add_paragraph("采购合同正文。")
        header = doc.sections[0].header
        header.is_linked_to_previous = False
        cls._add_payment_date(header.paragraphs[0])
        footer = doc.sections[0].footer
        footer.is_linked_to_previous = False
        cls._add_payment_date(footer.paragraphs[0])
        doc.save(path)

    def _round_trip(self, story: str) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / f"{story}-contract.docx"
            destination = root / f"{story}-contract-updated.docx"
            if story == "header":
                doc = Document()
                doc.add_paragraph("采购合同正文。")
                header = doc.sections[0].header
                header.is_linked_to_previous = False
                self._add_payment_date(header.paragraphs[0])
                doc.save(source)
            else:
                doc = Document()
                doc.add_paragraph("采购合同正文。")
                footer = doc.sections[0].footer
                footer.is_linked_to_previous = False
                self._add_payment_date(footer.paragraphs[0])
                doc.save(source)
            source_bytes = source.read_bytes()

            inspected = inspect_docx(source)
            self.assertTrue(inspected["ready"], inspected)
            self.assertEqual(inspected["payment_date_occurrence_count"], 1)
            occurrence = inspected["payment_date_occurrences"][0]
            self.assertTrue(
                str(occurrence["location"]).startswith(f"{story}:section:0/default/"),
                occurrence["location"],
            )
            self.assertEqual(occurrence["date"], "2026年9月15日")

            result = write_docx_copy(
                source,
                destination,
                replacement_date="2026年9月20日",
                precondition_identity=inspected["identity"],
            )

            # Source untouched, destination verified by the production path.
            self.assertEqual(source.read_bytes(), source_bytes)
            self.assertTrue(result["source_unchanged"])
            self.assertTrue(result["destination_reopened"])
            self.assertTrue(result["format_preserved"])
            self.assertTrue(result["surrounding_structure_preserved"])
            self.assertEqual(result["old_date"], "2026年9月15日")
            self.assertEqual(result["new_date"], "2026年9月20日")

            # Fresh reopen: the new date is exactly where the old one was.
            reopened = Document(destination)
            if story == "header":
                paragraph = reopened.sections[0].header.paragraphs[0]
            else:
                paragraph = reopened.sections[0].footer.paragraphs[0]
            self.assertEqual(paragraph.text, "付款日期：2026年9月20日")
            self.assertTrue(paragraph.runs[1].bold)
            self.assertAlmostEqual(paragraph.runs[1].font.size.pt, 11.0)
            self.assertTrue(paragraph.runs[2].italic)
            self.assertEqual(reopened.paragraphs[0].text, "采购合同正文。")

            final = inspect_docx(destination)
            self.assertTrue(final["ready"], final)
            self.assertEqual(final["payment_date_occurrence_count"], 1)
            self.assertEqual(
                final["payment_date_occurrences"][0]["date"], "2026年9月20日"
            )

    def test_header_payment_date_full_round_trip(self) -> None:
        self._round_trip("header")

    def test_footer_payment_date_full_round_trip(self) -> None:
        self._round_trip("footer")

    def test_duplicate_story_targets_stay_ambiguous(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "duplicate.docx"
            destination = root / "duplicate-updated.docx"
            self._fixture(source)

            inspected = inspect_docx(source)
            self.assertTrue(inspected["ready"], inspected)
            self.assertEqual(inspected["payment_date_occurrence_count"], 2)
            with self.assertRaises(ValueError):
                write_docx_copy(
                    source,
                    destination,
                    replacement_date="2026年9月20日",
                    precondition_identity=inspected["identity"],
                )
            self.assertFalse(destination.exists())


if __name__ == "__main__":
    unittest.main()

from __future__ import annotations

import json
import os
import tempfile
import unittest
from datetime import datetime, time, timedelta
from pathlib import Path

from docx import Document
from docx.shared import Pt

from zn_agent.core.provider_bridge import build_resident_runtime
from zn_agent.core.recovery_bounded_work import RecoveryBoundedWorkLedger
from zn_agent.core.work import WorkMessage

TASK = "找到我昨天下载的那份合同，把付款日期改成我们说好的日期，保存到项目文件夹。"
AGREED = "我们说好的付款日期是 2026年9月20日。"


class E2E09FindYesterdayDocxAndEditTests(unittest.TestCase):
    @staticmethod
    def _stamp(path: Path, days: int) -> None:
        now = datetime.now().astimezone()
        stamp = datetime.combine(now.date() + timedelta(days=days), time(12), tzinfo=now.tzinfo).timestamp()
        os.utime(path, (stamp, stamp))

    @staticmethod
    def _contract(path: Path, *, payment_dates=("2026年9月15日",)) -> None:
        doc = Document()
        doc.add_paragraph("采购合同")
        doc.add_paragraph("合同编号：ZN-2026-0915")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "甲方"
        table.cell(0, 1).text = "乙方"
        table.cell(1, 0).text = "A 公司"
        table.cell(1, 1).text = "B 公司"
        for payment_date in payment_dates:
            paragraph = doc.add_paragraph()
            paragraph.add_run("付款日期：")
            first = paragraph.add_run(payment_date[:5])
            first.bold = True
            first.font.size = Pt(11)
            second = paragraph.add_run(payment_date[5:])
            second.italic = True
        doc.add_paragraph("交付地点：项目现场。")
        doc.save(path)

    @staticmethod
    def _unrelated(path: Path) -> None:
        doc = Document()
        doc.add_paragraph("会议纪要")
        doc.add_paragraph("没有合同付款日期字段。")
        doc.save(path)

    @staticmethod
    def _setup(base: Path, project: Path, thread_id: str):
        resident = build_resident_runtime(config={"model": {}}, store_path=base / "kernel.db")
        ledger = RecoveryBoundedWorkLedger(resident)
        thread = ledger.create_thread(thread_id=thread_id)
        ledger.attach_workspace(thread.thread_id, project)
        ledger._append(
            thread,
            WorkMessage(
                message_id=f"fixture-{thread_id}",
                thread_id=thread.thread_id,
                role="user",
                text=AGREED,
            ),
        )
        return resident, ledger

    @staticmethod
    def _actions(resident, event_id: str):
        return [item for item in reversed(resident.body.recent_actions(512)) if item.event_id == event_id]

    @staticmethod
    def _run_to_terminal(resident, event_id: str):
        for _ in range(192):
            done = resident.result_for(event_id)
            if done is not None:
                return done
            result = resident.live_once()
            if result is not None and result.event.event_id == event_id:
                return result
        raise AssertionError(f"event did not terminate: {resident.store.get_working_state()}")

    def test_normal_language_closes_real_docx_copy_with_cross_run_format_and_fresh_reopen(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            downloads = base / "Downloads"
            project = base / "project"
            downloads.mkdir()
            project.mkdir()
            target = downloads / "contract-final.docx"
            old = downloads / "contract-old.docx"
            unrelated = downloads / "notes.docx"
            self._contract(target)
            self._contract(old)
            self._unrelated(unrelated)
            self._stamp(target, -1)
            self._stamp(unrelated, -1)
            self._stamp(old, -2)
            target_bytes = target.read_bytes()
            unrelated_bytes = unrelated.read_bytes()
            old_bytes = old.read_bytes()

            resident, ledger = self._setup(base, project, "e2e09")
            _, run = ledger.submit("e2e09", TASK, payload={"downloads_path": str(downloads)})
            self.assertTrue(run.success, run)
            self.assertEqual(run.model_invocations, 0)
            destination = project / "contract-final-updated.docx"
            self.assertTrue(destination.exists())
            self.assertEqual(target.read_bytes(), target_bytes)
            self.assertEqual(unrelated.read_bytes(), unrelated_bytes)
            self.assertEqual(old.read_bytes(), old_bytes)

            reopened = Document(destination)
            paragraph = next(value for value in reopened.paragraphs if "付款日期" in value.text)
            self.assertEqual(paragraph.text, "付款日期：2026年9月20日")
            self.assertTrue(paragraph.runs[1].bold)
            self.assertAlmostEqual(paragraph.runs[1].font.size.pt, 11.0)
            self.assertTrue(paragraph.runs[2].italic)
            self.assertEqual(paragraph.runs[2].text, "")
            self.assertIn("交付地点：项目现场。", [value.text for value in reopened.paragraphs])
            self.assertEqual(
                [cell.text for row in reopened.tables[0].rows for cell in row.cells],
                ["甲方", "乙方", "A 公司", "B 公司"],
            )

            actions = self._actions(resident, run.event.event_id)
            self.assertEqual(sum(item.kind == "write_docx_copy" for item in actions), 1)
            self.assertGreaterEqual(sum(item.kind == "inspect_docx" for item in actions), 4)
            items = ledger.list_work_items("e2e09", limit=64)
            evidence_item = next(item for item in items if "local_document_work:v1" in item.acceptance_criteria)
            evidence = json.loads(evidence_item.result)
            self.assertEqual(evidence["status"], "complete")
            self.assertTrue(evidence["verification"]["source_unchanged"])
            self.assertTrue(evidence["verification"]["payment_date_exact"])
            self.assertEqual(evidence["source_path"], str(target.resolve()))
            self.assertEqual(evidence["destination_path"], str(destination.resolve()))
            resident.store.close()

    def test_two_plausible_yesterday_contracts_stop_with_zero_mutation(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            downloads = base / "Downloads"
            project = base / "project"
            downloads.mkdir()
            project.mkdir()
            first = downloads / "contract-a.docx"
            second = downloads / "contract-b.docx"
            self._contract(first)
            self._contract(second)
            self._stamp(first, -1)
            self._stamp(second, -1)
            before = {first: first.read_bytes(), second: second.read_bytes()}
            resident, ledger = self._setup(base, project, "e2e09-ambiguous")
            _, run = ledger.submit("e2e09-ambiguous", TASK, payload={"downloads_path": str(downloads)})
            self.assertFalse(run.success)
            self.assertEqual(run.model_invocations, 0)
            self.assertIn("ambiguous_source", run.reason)
            self.assertFalse(list(project.glob("*.docx")))
            self.assertEqual(first.read_bytes(), before[first])
            self.assertEqual(second.read_bytes(), before[second])
            self.assertFalse([item for item in self._actions(resident, run.event.event_id) if item.kind == "write_docx_copy"])
            resident.store.close()

    def test_missing_agreed_date_and_multiple_payment_targets_stop_before_write(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            downloads = base / "Downloads"
            project = base / "project"
            downloads.mkdir()
            project.mkdir()
            target = downloads / "contract.docx"
            self._contract(target)
            self._stamp(target, -1)
            before = target.read_bytes()
            resident = build_resident_runtime(config={"model": {}}, store_path=base / "kernel.db")
            ledger = RecoveryBoundedWorkLedger(resident)
            ledger.create_thread(thread_id="missing")
            ledger.attach_workspace("missing", project)
            _, run = ledger.submit("missing", TASK, payload={"downloads_path": str(downloads)})
            self.assertFalse(run.success)
            self.assertIn("agreed_date_ambiguous", run.reason)
            self.assertEqual(target.read_bytes(), before)
            self.assertFalse(list(project.glob("*.docx")))
            resident.store.close()

        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            downloads = base / "Downloads"
            project = base / "project"
            downloads.mkdir()
            project.mkdir()
            target = downloads / "contract.docx"
            self._contract(target, payment_dates=("2026年9月15日", "2026年9月16日"))
            self._stamp(target, -1)
            before = target.read_bytes()
            resident, ledger = self._setup(base, project, "multi-target")
            _, run = ledger.submit("multi-target", TASK, payload={"downloads_path": str(downloads)})
            self.assertFalse(run.success)
            self.assertIn("ambiguous_payment_date_target", run.reason)
            self.assertEqual(target.read_bytes(), before)
            self.assertFalse(list(project.glob("*.docx")))
            resident.store.close()

    def test_source_drift_after_binding_rejects_stale_evidence_and_creates_no_copy(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            downloads = base / "Downloads"
            project = base / "project"
            downloads.mkdir()
            project.mkdir()
            target = downloads / "contract.docx"
            self._contract(target)
            self._stamp(target, -1)
            resident, ledger = self._setup(base, project, "e2e09-drift")
            _, event = ledger.start("e2e09-drift", TASK, payload={"downloads_path": str(downloads)})
            for _ in range(128):
                state = resident.store.get_working_state()
                if state.current_event_id == event.event_id and state.stage == "local_office_mutate":
                    break
                self.assertIsNone(resident.live_once())
            else:
                self.fail("E2E-09 did not reach the bound-before-mutation stage")
            doc = Document(target)
            doc.add_paragraph("外部刚刚更新的条款")
            doc.save(target)
            run = self._run_to_terminal(resident, event.event_id)
            self.assertFalse(run.success)
            self.assertIn("stale_source_evidence", run.reason)
            self.assertFalse((project / "contract-updated.docx").exists())
            self.assertTrue(any("外部刚刚更新的条款" in value.text for value in Document(target).paragraphs))
            resident.store.close()


if __name__ == "__main__":
    unittest.main()

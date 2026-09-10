from __future__ import annotations

import hashlib
import json
import os
import tempfile
import unittest
from datetime import datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference

from zn_agent.core.provider_bridge import build_resident_runtime
from zn_agent.core.recovery_bounded_work import RecoveryBoundedWorkLedger
from zn_agent.core.spreadsheet_work import AMOUNT_NUMBER_FORMAT

TASK = "把昨天那个表整理一下，重复项去掉，金额列统一格式，别动原文件，给我一个处理好的版本。"


class E2E10SpreadsheetCleanupTests(unittest.TestCase):
    @staticmethod
    def _stamp(path: Path, days: int) -> None:
        now = datetime.now().astimezone()
        stamp = datetime.combine(now.date() + timedelta(days=days), time(12), tzinfo=now.tzinfo).timestamp()
        os.utime(path, (stamp, stamp))

    @staticmethod
    def _sales(path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "销售"
        sheet.append(["订单号", "客户", "金额"])
        rows = [(1001, "A", 1200), (1002, "B", 99.5), (1001, "A", 1200), (1003, "C", 8), (1002, "B", 99.5)]
        formats = ["0", "0.0", "0.00", "General", "#,##0"]
        for index, row in enumerate(rows, start=2):
            sheet.append(row)
            sheet.cell(index, 3).number_format = formats[index - 2]
        workbook.save(path)

    @staticmethod
    def _setup(base: Path, project: Path, thread_id: str):
        resident = build_resident_runtime(config={"model": {}}, store_path=base / "kernel.db")
        ledger = RecoveryBoundedWorkLedger(resident)
        ledger.create_thread(thread_id=thread_id)
        ledger.attach_workspace(thread_id, project)
        return resident, ledger

    @staticmethod
    def _sha(path: Path) -> str:
        return hashlib.sha256(path.read_bytes()).hexdigest()

    @staticmethod
    def _actions(resident, event_id: str):
        return [item for item in reversed(resident.body.recent_actions(512)) if item.event_id == event_id]

    @staticmethod
    def _run_to_terminal(resident, event_id: str):
        for _ in range(192):
            done = resident.result_for(event_id)
            if done is not None:
                return done
            result = resident.live_once()
            if result is not None and result.event.event_id == event_id:
                return result
        raise AssertionError(f"event did not terminate: {resident.store.get_working_state()}")

    def test_normal_language_zero_model_cleanup_preserves_original_bytes_and_verifies_real_xlsx(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            downloads = base / "Downloads"
            project = base / "project"
            downloads.mkdir()
            project.mkdir()
            source = downloads / "sales.xlsx"
            old = downloads / "old-sales.xlsx"
            self._sales(source)
            self._sales(old)
            self._stamp(source, -1)
            self._stamp(old, -2)
            before_sha = self._sha(source)
            before_mtime = source.stat().st_mtime_ns

            resident, ledger = self._setup(base, project, "e2e10")
            _, run = ledger.submit("e2e10", TASK, payload={"downloads_path": str(downloads)})
            self.assertTrue(run.success, run)
            self.assertEqual(run.model_invocations, 0)
            destination = project / "sales-cleaned.xlsx"
            self.assertTrue(destination.exists())
            self.assertEqual(self._sha(source), before_sha)
            self.assertEqual(source.stat().st_mtime_ns, before_mtime)

            workbook = load_workbook(destination)
            sheet = workbook.active
            self.assertEqual(sheet.title, "销售")
            self.assertEqual(
                [tuple(cell.value for cell in row) for row in sheet.iter_rows()],
                [("订单号", "客户", "金额"), (1001, "A", 1200), (1002, "B", 99.5), (1003, "C", 8)],
            )
            self.assertTrue(all(sheet.cell(row, 3).data_type == "n" for row in range(2, sheet.max_row + 1)))
            self.assertEqual(
                [sheet.cell(row, 3).number_format for row in range(2, sheet.max_row + 1)],
                [AMOUNT_NUMBER_FORMAT] * 3,
            )
            actions = self._actions(resident, run.event.event_id)
            self.assertEqual(sum(item.kind == "write_xlsx_copy" for item in actions), 1)
            self.assertGreaterEqual(sum(item.kind == "inspect_xlsx" for item in actions), 3)
            items = ledger.list_work_items("e2e10", limit=64)
            evidence_item = next(item for item in items if "spreadsheet_cleanup_work:v1" in item.acceptance_criteria)
            evidence = json.loads(evidence_item.result)
            self.assertEqual(evidence["status"], "complete")
            self.assertTrue(evidence["verification"]["source_sha256_unchanged"])
            self.assertTrue(evidence["verification"]["zero_exact_duplicate_rows"])
            self.assertTrue(evidence["verification"]["amount_number_format_uniform"])
            resident.store.close()

    def test_two_yesterday_files_and_ambiguous_amount_columns_fail_closed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            downloads = base / "Downloads"
            project = base / "project"
            downloads.mkdir()
            project.mkdir()
            first = downloads / "a.xlsx"
            second = downloads / "b.xlsx"
            self._sales(first)
            self._sales(second)
            self._stamp(first, -1)
            self._stamp(second, -1)
            before = (self._sha(first), self._sha(second))
            resident, ledger = self._setup(base, project, "ambiguous-source")
            _, run = ledger.submit("ambiguous-source", TASK, payload={"downloads_path": str(downloads)})
            self.assertFalse(run.success)
            self.assertIn("ambiguous_source", run.reason)
            self.assertEqual((self._sha(first), self._sha(second)), before)
            self.assertFalse(list(project.glob("*.xlsx")))
            resident.store.close()

        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            downloads = base / "Downloads"
            project = base / "project"
            downloads.mkdir()
            project.mkdir()
            source = downloads / "amounts.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["订单号", "含税金额", "未税金额"])
            sheet.append([1, 10, 9])
            workbook.save(source)
            self._stamp(source, -1)
            before = self._sha(source)
            resident, ledger = self._setup(base, project, "ambiguous-amount")
            _, run = ledger.submit("ambiguous-amount", TASK, payload={"downloads_path": str(downloads)})
            self.assertFalse(run.success)
            self.assertIn("ambiguous_amount_column", run.reason)
            self.assertEqual(self._sha(source), before)
            self.assertFalse(list(project.glob("*.xlsx")))
            resident.store.close()

    def test_formula_and_complex_workbook_fail_closed(self) -> None:
        for mode in ("formula", "chart"):
            with self.subTest(mode=mode), tempfile.TemporaryDirectory() as tmp:
                base = Path(tmp)
                downloads = base / "Downloads"
                project = base / "project"
                downloads.mkdir()
                project.mkdir()
                source = downloads / f"{mode}.xlsx"
                workbook = Workbook()
                sheet = workbook.active
                sheet.append(["订单号", "金额"])
                sheet.append([1, 10])
                if mode == "formula":
                    sheet.cell(2, 2).value = "=5+5"
                else:
                    chart = BarChart()
                    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2))
                    sheet.add_chart(chart, "E2")
                workbook.save(source)
                self._stamp(source, -1)
                before = self._sha(source)
                resident, ledger = self._setup(base, project, f"blocked-{mode}")
                _, run = ledger.submit(f"blocked-{mode}", TASK, payload={"downloads_path": str(downloads)})
                self.assertFalse(run.success)
                self.assertIn("unsupported_workbook_structure", run.reason)
                self.assertEqual(self._sha(source), before)
                self.assertFalse(list(project.glob("*.xlsx")))
                resident.store.close()

    def test_source_drift_and_output_collision_stop_without_overwrite(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            downloads = base / "Downloads"
            project = base / "project"
            downloads.mkdir()
            project.mkdir()
            source = downloads / "sales.xlsx"
            self._sales(source)
            self._stamp(source, -1)
            resident, ledger = self._setup(base, project, "drift")
            _, event = ledger.start("drift", TASK, payload={"downloads_path": str(downloads)})
            for _ in range(128):
                state = resident.store.get_working_state()
                if state.current_event_id == event.event_id and state.stage == "local_office_mutate":
                    break
                self.assertIsNone(resident.live_once())
            else:
                self.fail("E2E-10 did not reach the bound-before-mutation stage")
            workbook = load_workbook(source)
            workbook.active.append([1004, "D", 5])
            workbook.save(source)
            run = self._run_to_terminal(resident, event.event_id)
            self.assertFalse(run.success)
            self.assertIn("stale_source_evidence", run.reason)
            self.assertFalse((project / "sales-cleaned.xlsx").exists())
            resident.store.close()

        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            downloads = base / "Downloads"
            project = base / "project"
            downloads.mkdir()
            project.mkdir()
            source = downloads / "sales.xlsx"
            self._sales(source)
            self._stamp(source, -1)
            before = self._sha(source)
            destination = project / "sales-cleaned.xlsx"
            destination.write_bytes(b"sentinel")
            resident, ledger = self._setup(base, project, "collision")
            _, run = ledger.submit("collision", TASK, payload={"downloads_path": str(downloads)})
            self.assertFalse(run.success)
            self.assertIn("output_collision", run.reason)
            self.assertEqual(destination.read_bytes(), b"sentinel")
            self.assertEqual(self._sha(source), before)
            self.assertFalse([item for item in self._actions(resident, run.event.event_id) if item.kind == "write_xlsx_copy"])
            resident.store.close()


if __name__ == "__main__":
    unittest.main()

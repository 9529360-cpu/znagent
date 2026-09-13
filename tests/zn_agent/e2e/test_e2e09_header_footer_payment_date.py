from __future__ import annotations

import os
import tempfile
import unittest
from datetime import datetime, time, timedelta
from pathlib import Path

from docx import Document
from docx.shared import Pt

from zn_agent.core.provider_bridge import build_resident_runtime
from zn_agent.core.work import WorkMessage

TASK = "找到我昨天下载的那份合同，把付款日期改成我们说好的日期，保存到项目文件夹。"
AGREED = "我们说好的付款日期是 2026年9月20日。"


class E2E09HeaderFooterPaymentDateTests(unittest.TestCase):
    @staticmethod
    def _stamp_yesterday(path: Path) -> None:
        now = datetime.now().astimezone()
        stamp = datetime.combine(
            now.date() + timedelta(days=-1),
            time(12),
            tzinfo=now.tzinfo,
        ).timestamp()
        os.utime(path, (stamp, stamp))

    @staticmethod
    def _contract_with_header_payment_date(path: Path) -> None:
        doc = Document()
        doc.add_paragraph("采购合同")
        doc.add_paragraph("合同编号：ZN-HF-2026")
        doc.add_paragraph("正文条款保持不变。")
        header = doc.sections[0].header
        header.is_linked_to_previous = False
        paragraph = header.paragraphs[0]
        paragraph.add_run("付款日期：")
        year = paragraph.add_run("2026年")
        year.bold = True
        year.font.size = Pt(11)
        month_day = paragraph.add_run("9月15日")
        month_day.italic = True
        doc.save(path)

    def test_normal_language_updates_unique_active_header_target_to_verified_copy(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            downloads = base / "Downloads"
            project = base / "project"
            downloads.mkdir()
            project.mkdir()
            source = downloads / "header-contract.docx"
            self._contract_with_header_payment_date(source)
            self._stamp_yesterday(source)
            source_bytes = source.read_bytes()

            db_path = base / "zn-e2e09-header-footer.db"
            resident = build_resident_runtime(config={"model": {}}, store_path=db_path)
            self.addCleanup(resident.store.close)
            ledger = resident.work_ledger
            thread = ledger.create_thread(thread_id="header-payment-date")
            thread = ledger.attach_workspace(thread.thread_id, project)
            ledger._append(
                thread,
                WorkMessage(
                    message_id="fixture-header-agreed-date",
                    thread_id=thread.thread_id,
                    role="user",
                    text=AGREED,
                ),
            )

            _, run = ledger.submit(
                thread.thread_id,
                TASK,
                payload={"downloads_path": str(downloads)},
            )
            self.assertTrue(run.success, run)
            self.assertEqual(run.model_invocations, 0)

            destination = project / "header-contract-updated.docx"
            self.assertTrue(destination.exists())
            self.assertEqual(source.read_bytes(), source_bytes)
            reopened = Document(destination)
            header = reopened.sections[0].header.paragraphs[0]
            self.assertEqual(header.text, "付款日期：2026年9月20日")
            self.assertTrue(header.runs[1].bold)
            self.assertAlmostEqual(header.runs[1].font.size.pt, 11.0)
            self.assertTrue(header.runs[2].italic)
            self.assertEqual(header.runs[2].text, "")
            self.assertIn(
                "正文条款保持不变。",
                [paragraph.text for paragraph in reopened.paragraphs],
            )

            actions = [
                item
                for item in reversed(resident.body.recent_actions(512))
                if item.event_id == run.event.event_id
            ]
            self.assertEqual(sum(item.kind == "write_docx_copy" for item in actions), 1)
            self.assertGreaterEqual(sum(item.kind == "inspect_docx" for item in actions), 4)


if __name__ == "__main__":
    unittest.main()

from __future__ import annotations

import hashlib
import http.server
import json
import os
import tempfile
import threading
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from zn_agent.core.browser import (
    BrowserAction,
    BrowserActionAuthority,
    BrowserActionKind,
    BrowserPermissionContext,
)
from zn_agent.core.provider_bridge import build_resident_runtime

TASK = "把这个网站里的数据整理进我现在这个表里。"


_TABLE = """
<table aria-label="Orders">
  <thead><tr><th>订单号</th><th>客户</th><th>金额</th></tr></thead>
  <tbody>
    <tr><td>1002</td><td>B</td><td id="drift-cell">99.50</td></tr>
    <tr><td>1003</td><td>C</td><td>8.00</td></tr>
  </tbody>
</table>
"""

_HIDDEN_ROW_TABLE = """
<table aria-label="Orders">
  <thead><tr><th>订单号</th><th>客户</th><th>金额</th></tr></thead>
  <tbody>
    <tr><td>1002</td><td>B</td><td>99.50</td></tr>
  </tbody>
  <tbody style="display:none">
    <tr><td>SECRET-1003</td><td>Hidden</td><td>8.00</td></tr>
  </tbody>
</table>
"""

_ARIA_HIDDEN_ROW_TABLE = """
<table aria-label="Orders">
  <thead><tr><th>订单号</th><th>客户</th><th>金额</th></tr></thead>
  <tbody>
    <tr><td>1002</td><td>B</td><td>99.50</td></tr>
  </tbody>
  <tbody aria-hidden="true">
    <tr><td>SECRET-1003</td><td>Hidden</td><td>8.00</td></tr>
  </tbody>
</table>
"""


class _FixtureHandler(http.server.BaseHTTPRequestHandler):
    def do_GET(self):
        if self.path == "/two":
            body = _TABLE + _TABLE.replace('aria-label="Orders"', 'aria-label="Other Orders"')
        elif self.path == "/hidden-row":
            body = _HIDDEN_ROW_TABLE
        elif self.path == "/aria-hidden-row":
            body = _ARIA_HIDDEN_ROW_TABLE
        else:
            body = _TABLE
        payload = (
            "<!doctype html><html><head><title>ZN E2E-11</title></head><body>"
            + body
            + "</body></html>"
        ).encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)

    def log_message(self, fmt, *args):
        return


class E2E11BrowserDataToSpreadsheetTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _FixtureHandler)
        cls.port = int(cls.server.server_address[1])
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.origin = f"http://127.0.0.1:{cls.port}"

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join(timeout=5)

    @staticmethod
    def _workbook(path: Path, *, formula: bool = False) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "销售"
        sheet.append(["订单号", "客户", "金额"])
        sheet.append([1001, "A", "=600+600" if formula else 1200])
        if not formula:
            sheet.cell(2, 3).number_format = "#,##0.00"
        workbook.save(path)

    @staticmethod
    def _sha(path: Path) -> str:
        return hashlib.sha256(path.read_bytes()).hexdigest()

    @staticmethod
    def _actions(resident, event_id: str):
        return [item for item in reversed(resident.body.recent_actions(512)) if item.event_id == event_id]

    @staticmethod
    def _run_to_terminal(resident, event_id: str):
        for _ in range(256):
            completed = resident.result_for(event_id)
            if completed is not None:
                return completed
            result = resident.live_once()
            if result is not None and result.event.event_id == event_id:
                return result
        raise AssertionError(f"event did not terminate: {resident.store.get_working_state()}")

    @staticmethod
    def _run_to_stage(resident, event_id: str, stage: str):
        for _ in range(160):
            state = resident.store.get_working_state()
            if state.current_event_id == event_id and state.stage == stage:
                return
            result = resident.live_once()
            if result is not None:
                raise AssertionError(f"event terminated before {stage}: {result}")
        raise AssertionError(f"event did not reach {stage}")

    def _setup(self, root: Path, thread_id: str, *, path: str = "/one"):
        db_tmp = tempfile.TemporaryDirectory()
        db_path = Path(db_tmp.name) / f"zn-e2e11-{thread_id}-{os.getpid()}.db"
        resident = build_resident_runtime(config={"model": {}}, store_path=db_path)
        self.addCleanup(db_tmp.cleanup)
        ledger = resident.work_ledger
        workspace = root / "workspace"
        workspace.mkdir()
        ledger.create_thread(thread_id=thread_id)
        ledger.attach_workspace(thread_id, workspace)

        permission = BrowserPermissionContext(
            allow_navigation=True,
            allow_page_interaction=True,
            allow_text_entry=False,
            allow_private_network=True,
            allowed_origins=(self.origin,),
        )
        browser = resident.managed_browser
        session = browser.open_session(permission=permission, headless=True)
        initial = browser.observe(session.session_id)
        url = self.origin + path
        action = BrowserAction.create(
            session_id=session.session_id,
            page_id=initial.page_id,
            kind=BrowserActionKind.NAVIGATE,
            args={"url": url},
            expected={"url_equals": url},
        )
        authority = BrowserActionAuthority.from_observation(action, initial, permission)
        effect = browser.act(action, authority)
        self.assertTrue(effect.success, effect.error)
        observed = browser.observe(session.session_id, page_id=initial.page_id)
        self.assertEqual(observed.url, url)

        def cleanup():
            try:
                browser.close_session(session.session_id)
            finally:
                resident.store.close()

        self.addCleanup(cleanup)
        return resident, ledger, workspace, session.session_id, observed.page_id

    @staticmethod
    def _payload(source: Path, session_id: str, page_id: str):
        return {
            "spreadsheet_path": str(source),
            "browser_session_id": session_id,
            "page_id": page_id,
        }

    def test_real_chromium_to_real_xlsx_happy_path_is_zero_model_and_fresh_verified(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            resident, ledger, workspace, session_id, page_id = self._setup(root, "happy")
            source = workspace / "sales.xlsx"
            self._workbook(source)
            before_sha = self._sha(source)
            before_mtime = source.stat().st_mtime_ns
            _, run = ledger.submit(
                "happy",
                TASK,
                payload=self._payload(source, session_id, page_id),
            )
            self.assertTrue(run.success, run)
            self.assertEqual(run.model_invocations, 0)
            self.assertEqual(self._sha(source), before_sha)
            self.assertEqual(source.stat().st_mtime_ns, before_mtime)

            destination = workspace / "sales-webdata.xlsx"
            self.assertTrue(destination.exists())
            workbook = load_workbook(destination)
            sheet = workbook.active
            self.assertEqual(sheet.title, "销售")
            self.assertEqual(
                [tuple(cell.value for cell in row) for row in sheet.iter_rows()],
                [
                    ("订单号", "客户", "金额"),
                    (1001, "A", 1200),
                    ("1002", "B", "99.50"),
                    ("1003", "C", "8.00"),
                ],
            )
            self.assertEqual(sheet.cell(2, 3).number_format, "#,##0.00")
            self.assertTrue(all(sheet.cell(row, 1).data_type == "s" for row in (3, 4)))

            actions = self._actions(resident, run.event.event_id)
            self.assertEqual(sum(item.kind == "append_xlsx_rows_copy" for item in actions), 1)
            self.assertGreaterEqual(sum(item.kind == "inspect_xlsx_append_target" for item in actions), 3)
            evidence_item = next(
                item
                for item in ledger.list_work_items("happy", limit=64)
                if "browser_spreadsheet_import:v1" in item.acceptance_criteria
            )
            evidence = json.loads(evidence_item.result)
            self.assertEqual(evidence["status"], "complete")
            self.assertEqual(evidence["imported_row_count"], 2)
            self.assertEqual(evidence["imported_column_count"], 3)
            self.assertEqual(
                evidence["fresh_browser_verification_fingerprint"],
                evidence["browser_source_fingerprint"],
            )
            self.assertTrue(evidence["fresh_destination_reopen_result"]["ready"])

    def test_real_browser_drift_after_capture_stops_before_any_destination(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            resident, ledger, workspace, session_id, page_id = self._setup(root, "browser-drift")
            source = workspace / "sales.xlsx"
            self._workbook(source)
            _, event = ledger.start(
                "browser-drift",
                TASK,
                payload=self._payload(source, session_id, page_id),
            )
            self._run_to_stage(resident, event.event_id, "browser_spreadsheet_mutate")
            page = resident.managed_browser._page(
                resident.managed_browser._session(session_id), page_id
            )
            page.evaluate("document.getElementById('drift-cell').textContent = '100.00'")
            run = self._run_to_terminal(resident, event.event_id)
            self.assertFalse(run.success)
            self.assertEqual(run.model_invocations, 0)
            self.assertIn("stale_browser_table_evidence", run.reason)
            self.assertFalse((workspace / "sales-webdata.xlsx").exists())
            self.assertFalse(
                any(item.kind == "append_xlsx_rows_copy" for item in self._actions(resident, event.event_id))
            )

    def test_real_hidden_table_region_fails_closed_before_xlsx_publish(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            resident, ledger, workspace, session_id, page_id = self._setup(
                root,
                "hidden-row",
                path="/hidden-row",
            )
            source = workspace / "sales.xlsx"
            self._workbook(source)
            _, run = ledger.submit(
                "hidden-row",
                TASK,
                payload=self._payload(source, session_id, page_id),
            )
            self.assertFalse(run.success)
            self.assertEqual(run.model_invocations, 0)
            self.assertIn("browser_table_observation_failed", run.reason)
            self.assertIn("hidden_row", run.reason)
            self.assertFalse((workspace / "sales-webdata.xlsx").exists())
            self.assertFalse(
                any(item.kind == "append_xlsx_rows_copy" for item in self._actions(resident, run.event.event_id))
            )

    def test_real_aria_hidden_ancestor_row_fails_closed_before_xlsx_publish(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            resident, ledger, workspace, session_id, page_id = self._setup(
                root,
                "aria-hidden-row",
                path="/aria-hidden-row",
            )
            source = workspace / "sales.xlsx"
            self._workbook(source)
            _, run = ledger.submit(
                "aria-hidden-row",
                TASK,
                payload=self._payload(source, session_id, page_id),
            )
            self.assertFalse(run.success)
            self.assertEqual(run.model_invocations, 0)
            self.assertIn("browser_table_observation_failed", run.reason)
            self.assertIn("hidden_row", run.reason)
            self.assertFalse((workspace / "sales-webdata.xlsx").exists())
            self.assertFalse(
                any(item.kind == "append_xlsx_rows_copy" for item in self._actions(resident, run.event.event_id))
            )

    def test_real_multiple_tables_fail_closed(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            resident, ledger, workspace, session_id, page_id = self._setup(
                root,
                "two-tables",
                path="/two",
            )
            source = workspace / "sales.xlsx"
            self._workbook(source)
            _, run = ledger.submit(
                "two-tables",
                TASK,
                payload=self._payload(source, session_id, page_id),
            )
            self.assertFalse(run.success)
            self.assertIn("ambiguous_browser_table", run.reason)
            self.assertFalse((workspace / "sales-webdata.xlsx").exists())

    def test_real_output_collision_fails_closed(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            resident, ledger, workspace, session_id, page_id = self._setup(root, "collision")
            source = workspace / "sales.xlsx"
            self._workbook(source)
            destination = workspace / "sales-webdata.xlsx"
            destination.write_bytes(b"sentinel")
            _, run = ledger.submit(
                "collision",
                TASK,
                payload=self._payload(source, session_id, page_id),
            )
            self.assertFalse(run.success)
            self.assertIn("output_collision", run.reason)
            self.assertEqual(destination.read_bytes(), b"sentinel")

    def test_real_complex_xlsx_formula_fails_closed(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            resident, ledger, workspace, session_id, page_id = self._setup(root, "formula")
            source = workspace / "sales.xlsx"
            self._workbook(source, formula=True)
            _, run = ledger.submit(
                "formula",
                TASK,
                payload=self._payload(source, session_id, page_id),
            )
            self.assertFalse(run.success)
            self.assertIn("unsupported_workbook_structure", run.reason)
            self.assertFalse((workspace / "sales-webdata.xlsx").exists())


if __name__ == "__main__":
    unittest.main(verbosity=2)

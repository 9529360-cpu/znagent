from __future__ import annotations

"""Deterministic bounded DOCX inspection and run-aware copy mutation.

The model is intentionally absent from this module. It accepts already-bound
source identity and explicit deterministic replacement authority, mutates only
ordinary WordprocessingML body/table-cell run spans plus explicitly active
header/footer stories, reopens the package, and proves source/destination
invariants before reporting success.
"""

import hashlib
import json
import os
import re
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Iterator, Mapping
from xml.etree import ElementTree as ET

from docx import Document
from docx.document import Document as _Document
from docx.text.paragraph import Paragraph

from .file_identity import compare_file_identities, observe_file_identity

MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_PARAGRAPHS = 512
MAX_VISIBLE_CHARS = 128 * 1024
MAX_COMPLETION_TARGETS = 3
MAX_COMPLETION_QUESTION_CHARS = 300
MAX_COMPLETION_CONTEXT_CHARS = 800
MAX_COMPLETION_REPLACEMENT_CHARS = 2_000

_DATE = r"(?:\d{4}年\d{1,2}月\d{1,2}日|\d{4}[-/.]\d{1,2}[-/.]\d{1,2})"
PAYMENT_DATE_RE = re.compile(rf"(?P<label>付款日期\s*[:：]?\s*)(?P<date>{_DATE})")
COMPLETION_PLACEHOLDER_RE = re.compile(r"【待补充：(?P<question>[^】]*)】")
_COMPLETION_PREFIX = "【待补充"

_UNSUPPORTED_PART_PREFIXES = (
    "word/embeddings/",
    "word/activeX/",
    "word/diagrams/",
    "word/charts/",
)
_UNSUPPORTED_PARTS = {
    "word/vbaProject.bin",
    "word/comments.xml",
    "word/footnotes.xml",
    "word/endnotes.xml",
}
_UNSUPPORTED_XML_LOCALNAMES = {
    "altChunk",
    "commentRangeStart",
    "commentRangeEnd",
    "commentReference",
    "del",
    "drawing",
    "fldChar",
    "fldSimple",
    "gridSpan",
    "ins",
    "instrText",
    "object",
    "pict",
    "sdt",
    "smartTag",
    "txbxContent",
    "vMerge",
}


def _block(identity: dict[str, Any], detail: str) -> dict[str, Any]:
    return {
        "ready": False,
        "blocker": "unsupported_document_structure",
        "detail": str(detail)[:1000],
        "identity": identity,
    }


def _completion_block(
    identity: dict[str, Any],
    blocker: str,
    detail: str,
    *,
    target_count: int | None = None,
) -> dict[str, Any]:
    result: dict[str, Any] = {
        "ready": False,
        "blocker": str(blocker or "unsupported_document_structure"),
        "detail": str(detail)[:1000],
        "identity": identity,
        "target_count": target_count,
        "targets": [],
    }
    return result


def _localname(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag.rsplit(":", 1)[-1]


def _xml_root(payload: bytes) -> ET.Element:
    return ET.fromstring(payload)


def _xml_visible_text(payload: bytes) -> str:
    root = _xml_root(payload)
    chunks: list[str] = []
    for element in root.iter():
        if _localname(element.tag) in {"t", "tab", "br", "cr"}:
            if _localname(element.tag) == "t":
                chunks.append(element.text or "")
            elif _localname(element.tag) == "tab":
                chunks.append("\t")
            else:
                chunks.append("\n")
    return "".join(chunks)


def _xml_payment_occurrence_count(payload: bytes) -> int:
    """Count payment targets per WordprocessingML paragraph without story joins."""
    root = _xml_root(payload)
    count = 0
    for paragraph in root.iter():
        if _localname(paragraph.tag) != "p":
            continue
        chunks: list[str] = []
        for element in paragraph.iter():
            local = _localname(element.tag)
            if local == "t":
                chunks.append(element.text or "")
            elif local == "tab":
                chunks.append("\t")
            elif local in {"br", "cr"}:
                chunks.append("\n")
        count += len(list(PAYMENT_DATE_RE.finditer("".join(chunks))))
    return count


def _preflight(path: Path, identity: dict[str, Any]) -> tuple[bool, str | None]:
    if path.suffix.casefold() != ".docx":
        return False, "only standard .docx is supported"
    if (
        identity.get("exists") is not True
        or identity.get("type") != "file"
        or identity.get("stable") is not True
        or identity.get("digest_complete") is not True
    ):
        return False, "source file identity is not a stable complete regular-file observation"
    if int(identity.get("size_bytes") or 0) > MAX_FILE_BYTES:
        return False, "DOCX exceeds the bounded file-size limit"
    if not zipfile.is_zipfile(path):
        return False, "encrypted, legacy, corrupt, or non-OOXML Word packages are unsupported"
    try:
        with zipfile.ZipFile(path, "r") as package:
            names = set(package.namelist())
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                return False, "required WordprocessingML package parts are missing"
            for name in names:
                if name in _UNSUPPORTED_PARTS or any(
                    name.startswith(prefix) for prefix in _UNSUPPORTED_PART_PREFIXES
                ):
                    return False, f"unsupported DOCX package part: {name}"
            main_xml = package.read("word/document.xml")
            root = _xml_root(main_xml)
            bad = sorted(
                {_localname(el.tag) for el in root.iter()}
                & _UNSUPPORTED_XML_LOCALNAMES
            )
            if bad:
                return False, "unsupported WordprocessingML structures: " + ", ".join(bad)
    except (OSError, KeyError, zipfile.BadZipFile, ET.ParseError) as exc:
        return False, f"DOCX package preflight failed: {type(exc).__name__}: {exc}"
    return True, None


def _header_footer_payment_story_info(path: Path) -> tuple[int, str | None]:
    """Return bounded header/footer target count and reject complex target stories.

    Existing body-only documents keep their prior behavior. Header/footer XML is
    admitted only when the target-bearing story itself contains ordinary text
    structures that python-docx can round-trip without us claiming authority over
    fields, drawings, content controls, tracked changes, or text boxes.
    """
    count = 0
    try:
        with zipfile.ZipFile(path, "r") as package:
            for name in sorted(package.namelist()):
                if not (
                    name.startswith("word/header") or name.startswith("word/footer")
                ) or not name.endswith(".xml"):
                    continue
                payload = package.read(name)
                story_count = _xml_payment_occurrence_count(payload)
                if not story_count:
                    continue
                root = _xml_root(payload)
                bad = sorted(
                    {_localname(el.tag) for el in root.iter()}
                    & _UNSUPPORTED_XML_LOCALNAMES
                )
                if bad:
                    return count, (
                        f"payment-date header/footer story {name} contains unsupported "
                        "WordprocessingML structures: " + ", ".join(bad)
                    )
                count += story_count
    except (OSError, KeyError, zipfile.BadZipFile, ET.ParseError) as exc:
        return count, (
            "DOCX header/footer payment-story inspection failed: "
            f"{type(exc).__name__}: {exc}"
        )
    return count, None


def _completion_story_blocker(path: Path) -> str | None:
    """Reject completion authority outside the supported body/table story."""
    try:
        with zipfile.ZipFile(path, "r") as package:
            for name in sorted(package.namelist()):
                if not (
                    name.startswith("word/header") or name.startswith("word/footer")
                ) or not name.endswith(".xml"):
                    continue
                if _COMPLETION_PREFIX in _xml_visible_text(package.read(name)):
                    return "completion placeholder occurs in a header/footer story"
    except (OSError, KeyError, zipfile.BadZipFile, ET.ParseError) as exc:
        return f"DOCX completion story inspection failed: {type(exc).__name__}: {exc}"
    return None


def _iter_container_paragraphs(container: Any, prefix: str) -> Iterator[tuple[str, Paragraph]]:
    for index, paragraph in enumerate(container.paragraphs):
        yield f"{prefix}/p:{index}", paragraph
    for table_index, table in enumerate(container.tables):
        for row_index, row in enumerate(table.rows):
            for col_index, cell in enumerate(row.cells):
                if cell.tables:
                    raise ValueError("nested tables are outside DOCX v1 scope")
                for paragraph_index, paragraph in enumerate(cell.paragraphs):
                    yield (
                        f"{prefix}/table:{table_index}/cell:{row_index},{col_index}/p:{paragraph_index}",
                        paragraph,
                    )


def _iter_active_header_footer_paragraphs(
    document: _Document,
) -> Iterator[tuple[str, Paragraph]]:
    odd_even = bool(document.settings.odd_and_even_pages_header_footer)
    for section_index, section in enumerate(document.sections):
        stories = (
            ("header", "default", section.header, True),
            ("footer", "default", section.footer, True),
            (
                "header",
                "first",
                section.first_page_header,
                bool(section.different_first_page_header_footer),
            ),
            (
                "footer",
                "first",
                section.first_page_footer,
                bool(section.different_first_page_header_footer),
            ),
            ("header", "even", section.even_page_header, odd_even),
            ("footer", "even", section.even_page_footer, odd_even),
        )
        for kind, variant, story, active in stories:
            # Accessing paragraphs on a linked story can create a definition. Check
            # linkage first and only touch stories with an explicit active owner.
            if not active or story.is_linked_to_previous:
                continue
            prefix = f"{kind}:section:{section_index}/{variant}"
            yield from _iter_container_paragraphs(story, prefix)


def _iter_paragraphs(
    document: _Document,
    *,
    include_header_footer: bool = False,
) -> Iterator[tuple[str, Paragraph]]:
    yield from _iter_container_paragraphs(document, "body")
    if include_header_footer:
        yield from _iter_active_header_footer_paragraphs(document)


def _run_format(run) -> dict[str, Any]:
    size = run.font.size
    return {
        "bold": run.bold,
        "italic": run.italic,
        "underline": str(run.underline) if run.underline is not None else None,
        "style": getattr(run.style, "name", None),
        "font_name": run.font.name,
        "font_size_pt": float(size.pt) if size is not None else None,
    }


def _records(
    document: _Document,
    *,
    include_header_footer: bool = False,
) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    total_chars = 0
    for location, paragraph in _iter_paragraphs(
        document,
        include_header_footer=include_header_footer,
    ):
        if len(records) >= MAX_PARAGRAPHS:
            raise ValueError("document exceeds bounded paragraph count")
        text = paragraph.text
        total_chars += len(text)
        if total_chars > MAX_VISIBLE_CHARS:
            raise ValueError("document exceeds bounded visible-text limit")
        records.append(
            {
                "location": location,
                "text": text,
                "run_text": [run.text for run in paragraph.runs],
                "run_format": [_run_format(run) for run in paragraph.runs],
            }
        )
    return records


def _structure_fingerprint(records: list[dict[str, Any]]) -> str:
    payload = [
        {
            "location": record["location"],
            "text": record["text"],
            "run_format": record["run_format"],
        }
        for record in records
    ]
    return hashlib.sha256(
        json.dumps(
            payload,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()


def _target_occurrences(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    occurrences: list[dict[str, Any]] = []
    for record in records:
        text = str(record["text"])
        run_visible = "".join(str(value) for value in record["run_text"])
        for match in PAYMENT_DATE_RE.finditer(text):
            occurrences.append(
                {
                    "location": record["location"],
                    "date": match.group("date"),
                    "date_start": match.start("date"),
                    "date_end": match.end("date"),
                    "paragraph_sha256": hashlib.sha256(
                        text.encode("utf-8")
                    ).hexdigest(),
                    "run_mappable": run_visible == text,
                }
            )
    return occurrences


def _completion_context(text: str, start: int, end: int) -> str:
    budget = MAX_COMPLETION_CONTEXT_CHARS
    if len(text) <= budget:
        return text
    left_budget = budget // 2
    right_budget = budget - left_budget
    left = max(0, start - left_budget)
    right = min(len(text), end + right_budget)
    if right - left < budget:
        if left == 0:
            right = min(len(text), budget)
        elif right == len(text):
            left = max(0, len(text) - budget)
    return text[left:right]


def _completion_target_id(
    *,
    location: str,
    span_start: int,
    span_end: int,
    placeholder_text: str,
    paragraph_sha256: str,
) -> str:
    payload = {
        "location": location,
        "span_start": span_start,
        "span_end": span_end,
        "placeholder_sha256": hashlib.sha256(
            placeholder_text.encode("utf-8")
        ).hexdigest(),
        "paragraph_sha256": paragraph_sha256,
    }
    digest = hashlib.sha256(
        json.dumps(
            payload,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()
    return f"doc-target-{digest[:24]}"


def _completion_targets(
    records: list[dict[str, Any]],
    *,
    enforce_max: bool = True,
) -> list[dict[str, Any]]:
    targets: list[dict[str, Any]] = []
    for record in records:
        text = str(record["text"])
        matches = list(COMPLETION_PLACEHOLDER_RE.finditer(text))
        if text.count(_COMPLETION_PREFIX) != len(matches):
            raise ValueError(
                "malformed completion placeholder; expected exact 【待补充：<question>】 syntax"
            )
        run_visible = "".join(str(value) for value in record["run_text"])
        paragraph_sha256 = hashlib.sha256(text.encode("utf-8")).hexdigest()
        for match in matches:
            question = str(match.group("question") or "").strip()
            if not question:
                raise ValueError("completion placeholder question must not be empty")
            if len(question) > MAX_COMPLETION_QUESTION_CHARS:
                raise ValueError(
                    f"completion placeholder question exceeds {MAX_COMPLETION_QUESTION_CHARS} characters"
                )
            placeholder_text = match.group(0)
            span_start = match.start()
            span_end = match.end()
            targets.append(
                {
                    "target_id": _completion_target_id(
                        location=str(record["location"]),
                        span_start=span_start,
                        span_end=span_end,
                        placeholder_text=placeholder_text,
                        paragraph_sha256=paragraph_sha256,
                    ),
                    "location": record["location"],
                    "placeholder_text": placeholder_text,
                    "question": question,
                    "span_start": span_start,
                    "span_end": span_end,
                    "paragraph_sha256": paragraph_sha256,
                    "context_text": _completion_context(text, span_start, span_end),
                    "run_mappable": run_visible == text,
                }
            )
            if enforce_max and len(targets) > MAX_COMPLETION_TARGETS:
                raise OverflowError(
                    f"completion target count exceeds {MAX_COMPLETION_TARGETS}"
                )
    return targets


def inspect_docx(path: str | Path) -> dict[str, Any]:
    source = Path(path).resolve(strict=False)
    identity = observe_file_identity(source, max_hash_bytes=MAX_FILE_BYTES)
    ok, detail = _preflight(source, identity)
    if not ok:
        return _block(identity, detail or "unsupported DOCX")
    header_footer_target_count, story_detail = _header_footer_payment_story_info(source)
    if story_detail:
        return _block(identity, story_detail)
    try:
        document = Document(source)
        records = _records(
            document,
            include_header_footer=header_footer_target_count > 0,
        )
        occurrences = _target_occurrences(records)
    except Exception as exc:
        return _block(
            identity,
            f"DOCX parse/structure inspection failed: {type(exc).__name__}: {exc}",
        )
    supported_story_count = sum(
        str(item["location"]).startswith(("header:", "footer:"))
        for item in occurrences
    )
    if supported_story_count != header_footer_target_count:
        return _block(
            identity,
            "payment-date target occurs in an inactive, unowned, or unsupported header/footer story",
        )
    if any(not item["run_mappable"] for item in occurrences):
        return _block(
            identity,
            "payment-date target crosses unsupported paragraph children such as hyperlinks",
        )
    return {
        "ready": True,
        "blocker": None,
        "identity": identity,
        "paragraph_count": len(records),
        "table_count": len(document.tables),
        "payment_date_occurrence_count": len(occurrences),
        "payment_date_occurrences": occurrences,
        "structure_fingerprint": _structure_fingerprint(records),
    }


def inspect_docx_completion_targets(path: str | Path) -> dict[str, Any]:
    """Inspect explicit completion placeholders without granting mutation authority.

    Zero targets is a valid observation so the same primitive can independently
    prove a completed output has no remaining placeholder. Product admission is
    responsible for requiring 1..3 targets on the source document.
    """

    source = Path(path).resolve(strict=False)
    identity = observe_file_identity(source, max_hash_bytes=MAX_FILE_BYTES)
    ok, detail = _preflight(source, identity)
    if not ok:
        return _completion_block(
            identity,
            "unsupported_document_structure",
            detail or "unsupported DOCX",
        )
    story_blocker = _completion_story_blocker(source)
    if story_blocker:
        return _completion_block(
            identity,
            "unsupported_document_structure",
            story_blocker,
        )
    try:
        document = Document(source)
        records = _records(document)
        targets = _completion_targets(records)
    except OverflowError as exc:
        return _completion_block(
            identity,
            "completion_target_count_out_of_bounds",
            str(exc),
            target_count=MAX_COMPLETION_TARGETS + 1,
        )
    except ValueError as exc:
        return _completion_block(
            identity,
            "invalid_completion_placeholder",
            str(exc),
        )
    except Exception as exc:
        return _completion_block(
            identity,
            "unsupported_document_structure",
            f"DOCX completion inspection failed: {type(exc).__name__}: {exc}",
        )
    if any(not item["run_mappable"] for item in targets):
        return _completion_block(
            identity,
            "unsupported_document_structure",
            "completion target crosses unsupported paragraph children such as hyperlinks",
            target_count=len(targets),
        )
    return {
        "ready": True,
        "blocker": None,
        "identity": identity,
        "paragraph_count": len(records),
        "table_count": len(document.tables),
        "target_count": len(targets),
        "targets": targets,
        "structure_fingerprint": _structure_fingerprint(records),
    }


def replace_visible_span_across_runs(
    paragraph: Paragraph,
    start: int,
    end: int,
    replacement: str,
) -> None:
    """Replace one visible-text span without rebuilding the paragraph."""
    if start < 0 or end <= start:
        raise ValueError("replacement span is invalid")
    runs = list(paragraph.runs)
    visible = "".join(run.text for run in runs)
    if visible != paragraph.text:
        raise ValueError(
            "paragraph contains unsupported visible children for run-aware replacement"
        )
    if end > len(visible):
        raise ValueError("replacement span exceeds paragraph text")

    cursor = 0
    inserted = False
    touched = False
    for run in runs:
        text = run.text
        run_start, run_end = cursor, cursor + len(text)
        cursor = run_end
        overlap_start, overlap_end = max(start, run_start), min(end, run_end)
        if overlap_start >= overlap_end:
            continue
        touched = True
        before = text[: max(0, start - run_start)] if run_start < start else ""
        after = text[max(0, end - run_start) :] if run_end > end else ""
        middle = replacement if not inserted else ""
        run.text = before + middle + after
        inserted = True
    if not touched or not inserted:
        raise ValueError("replacement span did not map to ordinary text runs")


def _find_paragraph(document: _Document, location: str) -> Paragraph:
    include_header_footer = str(location).startswith(("header:", "footer:"))
    for current, paragraph in _iter_paragraphs(
        document,
        include_header_footer=include_header_footer,
    ):
        if current == location:
            return paragraph
    raise ValueError("bound DOCX target location no longer exists")


def _publish_no_overwrite(temp_path: Path, destination: Path) -> None:
    try:
        os.link(temp_path, destination)
    except FileExistsError:
        raise FileExistsError("destination already exists; refusing overwrite")
    temp_path.unlink()


def _normalized_completion_replacements(
    replacements: Any,
) -> dict[str, str]:
    rows: list[tuple[Any, Any]] = []
    if isinstance(replacements, Mapping):
        rows = list(replacements.items())
    elif isinstance(replacements, (list, tuple)):
        for item in replacements:
            if not isinstance(item, Mapping):
                raise ValueError("each completion replacement must be an object")
            if set(item.keys()) != {"target_id", "text"}:
                raise ValueError(
                    "completion replacement objects allow only target_id and text"
                )
            rows.append((item.get("target_id"), item.get("text")))
    else:
        raise ValueError("completion replacements must be a mapping or list")

    normalized: dict[str, str] = {}
    for raw_target, raw_text in rows:
        target_id = str(raw_target or "").strip()
        if not target_id:
            raise ValueError("completion replacement target_id must not be empty")
        if target_id in normalized:
            raise ValueError("completion replacement target_id must be unique")
        if not isinstance(raw_text, str):
            raise ValueError("completion replacement text must be a string")
        text = raw_text.strip()
        if not text:
            raise ValueError("completion replacement text must not be empty")
        if len(text) > MAX_COMPLETION_REPLACEMENT_CHARS:
            raise ValueError(
                f"completion replacement exceeds {MAX_COMPLETION_REPLACEMENT_CHARS} characters"
            )
        if _COMPLETION_PREFIX in text:
            raise ValueError("completion replacement must not create another placeholder")
        normalized[target_id] = text
    return normalized


def write_docx_copy(
    source_path: str | Path,
    destination_path: str | Path,
    *,
    replacement_date: str,
    precondition_identity: dict[str, Any],
) -> dict[str, Any]:
    source = Path(source_path).resolve(strict=True)
    destination = Path(destination_path).resolve(strict=False)
    if source == destination:
        raise ValueError("DOCX copy destination must differ from source")
    if destination.exists():
        raise FileExistsError("destination already exists; refusing overwrite")
    if not destination.parent.is_dir():
        raise ValueError("destination parent is unavailable")

    current = observe_file_identity(source, max_hash_bytes=MAX_FILE_BYTES)
    if compare_file_identities(precondition_identity, current).get("exact") is not True:
        raise RuntimeError(
            "stale_source_evidence: exact DOCX identity changed before mutation"
        )
    inspection = inspect_docx(source)
    if inspection.get("ready") is not True:
        raise RuntimeError(
            f"{inspection.get('blocker')}: {inspection.get('detail')}"
        )
    occurrences = list(inspection.get("payment_date_occurrences") or ())
    if len(occurrences) != 1:
        raise ValueError(
            "DOCX mutation requires exactly one payment-date occurrence"
        )
    target = occurrences[0]
    old_date = str(target["date"])
    if old_date == replacement_date:
        raise ValueError("replacement date already equals the current payment date")

    include_header_footer = str(target["location"]).startswith(("header:", "footer:"))
    document = Document(source)
    before_records = _records(
        document,
        include_header_footer=include_header_footer,
    )
    paragraph = _find_paragraph(document, str(target["location"]))
    before_target_text = paragraph.text
    if (
        hashlib.sha256(before_target_text.encode("utf-8")).hexdigest()
        != target["paragraph_sha256"]
    ):
        raise RuntimeError(
            "stale_source_evidence: bound paragraph changed before mutation"
        )
    replace_visible_span_across_runs(
        paragraph,
        int(target["date_start"]),
        int(target["date_end"]),
        replacement_date,
    )
    expected_target_text = (
        before_target_text[: int(target["date_start"])]
        + replacement_date
        + before_target_text[int(target["date_end"]) :]
    )

    fd, raw_temp = tempfile.mkstemp(
        prefix=f".{destination.name}.",
        suffix=".tmp",
        dir=destination.parent,
    )
    os.close(fd)
    temp_path = Path(raw_temp)
    try:
        document.save(temp_path)
        temp_document = Document(temp_path)
        after_records = _records(
            temp_document,
            include_header_footer=include_header_footer,
        )
        if len(after_records) != len(before_records):
            raise RuntimeError(
                "DOCX verification failed: paragraph topology changed"
            )
        changed = 0
        for before, after in zip(before_records, after_records):
            if before["location"] != after["location"]:
                raise RuntimeError(
                    "DOCX verification failed: paragraph location order changed"
                )
            if before["run_format"] != after["run_format"]:
                raise RuntimeError("DOCX verification failed: run formatting changed")
            expected_text = (
                expected_target_text
                if before["location"] == target["location"]
                else before["text"]
            )
            if after["text"] != expected_text:
                raise RuntimeError(
                    "DOCX verification failed: unrelated or target paragraph text changed unexpectedly"
                )
            changed += int(before["text"] != after["text"])
        if changed != 1:
            raise RuntimeError(
                "DOCX verification failed: expected exactly one paragraph text change"
            )
        temp_occurrences = _target_occurrences(after_records)
        if (
            len(temp_occurrences) != 1
            or temp_occurrences[0]["date"] != replacement_date
        ):
            raise RuntimeError(
                "DOCX verification failed: replacement date was not uniquely re-parsed"
            )

        source_after_temp = observe_file_identity(
            source,
            max_hash_bytes=MAX_FILE_BYTES,
        )
        if (
            compare_file_identities(current, source_after_temp).get("exact")
            is not True
        ):
            raise RuntimeError(
                "DOCX source changed while destination was being prepared"
            )
        if destination.exists():
            raise FileExistsError(
                "destination appeared during mutation; refusing overwrite"
            )
        _publish_no_overwrite(temp_path, destination)

        final_inspection = inspect_docx(destination)
        destination_identity = final_inspection.get("identity") or {}
        source_after = observe_file_identity(source, max_hash_bytes=MAX_FILE_BYTES)
        if compare_file_identities(current, source_after).get("exact") is not True:
            raise RuntimeError("DOCX source changed before final verification")
        if final_inspection.get("ready") is not True:
            raise RuntimeError("DOCX final reopen/inspection failed")
        final_occurrences = final_inspection.get("payment_date_occurrences") or []
        if (
            len(final_occurrences) != 1
            or final_occurrences[0].get("date") != replacement_date
        ):
            raise RuntimeError(
                "DOCX final verification did not find exactly the expected payment date"
            )
        if final_inspection.get("structure_fingerprint") != _structure_fingerprint(
            after_records
        ):
            raise RuntimeError(
                "DOCX final structure fingerprint changed after publication"
            )
        return {
            "source_identity_before": current,
            "source_identity_after": source_after,
            "destination_identity": destination_identity,
            "source_unchanged": True,
            "destination_reopened": True,
            "old_date": old_date,
            "new_date": replacement_date,
            "target_location": target["location"],
            "target_count_before": 1,
            "target_count_after": 1,
            "format_preserved": True,
            "surrounding_structure_preserved": True,
            "expected_structure_fingerprint": final_inspection[
                "structure_fingerprint"
            ],
        }
    except Exception:
        if temp_path.exists():
            temp_path.unlink(missing_ok=True)
        raise


def write_docx_completion_copy(
    source_path: str | Path,
    destination_path: str | Path,
    *,
    replacements: Any,
    precondition_identity: dict[str, Any],
) -> dict[str, Any]:
    """Write a verified no-overwrite DOCX copy for exact bound placeholders."""

    source = Path(source_path).resolve(strict=True)
    destination = Path(destination_path).resolve(strict=False)
    if source == destination:
        raise ValueError("DOCX completion destination must differ from source")
    if destination.exists():
        raise FileExistsError("destination already exists; refusing overwrite")
    if not destination.parent.is_dir():
        raise ValueError("destination parent is unavailable")

    current = observe_file_identity(source, max_hash_bytes=MAX_FILE_BYTES)
    if compare_file_identities(precondition_identity, current).get("exact") is not True:
        raise RuntimeError(
            "stale_source_evidence: exact DOCX identity changed before completion mutation"
        )

    inspection = inspect_docx_completion_targets(source)
    if inspection.get("ready") is not True:
        raise RuntimeError(
            f"{inspection.get('blocker')}: {inspection.get('detail')}"
        )
    targets = list(inspection.get("targets") or ())
    if not 1 <= len(targets) <= MAX_COMPLETION_TARGETS:
        raise ValueError(
            "DOCX completion mutation requires between one and three bound targets"
        )
    normalized = _normalized_completion_replacements(replacements)
    expected_ids = [str(target["target_id"]) for target in targets]
    if len(normalized) != len(expected_ids) or set(normalized) != set(expected_ids):
        raise RuntimeError(
            "target_drift: replacement target set does not exactly match fresh DOCX targets"
        )

    document = Document(source)
    before_records = _records(document)
    if _structure_fingerprint(before_records) != inspection.get(
        "structure_fingerprint"
    ):
        raise RuntimeError(
            "stale_source_evidence: source structure changed after target inspection"
        )

    targets_by_location: dict[str, list[dict[str, Any]]] = {}
    for target in targets:
        targets_by_location.setdefault(str(target["location"]), []).append(target)

    expected_text_by_location: dict[str, str] = {}
    target_verification: list[dict[str, Any]] = []
    for location, location_targets in targets_by_location.items():
        paragraph = _find_paragraph(document, location)
        original_text = paragraph.text
        expected_paragraph_hashes = {
            str(target["paragraph_sha256"]) for target in location_targets
        }
        if len(expected_paragraph_hashes) != 1 or hashlib.sha256(
            original_text.encode("utf-8")
        ).hexdigest() not in expected_paragraph_hashes:
            raise RuntimeError(
                "target_drift: bound completion paragraph changed before mutation"
            )
        expected_text = original_text
        for target in sorted(
            location_targets,
            key=lambda value: int(value["span_start"]),
            reverse=True,
        ):
            target_id = str(target["target_id"])
            start = int(target["span_start"])
            end = int(target["span_end"])
            placeholder_text = str(target["placeholder_text"])
            if original_text[start:end] != placeholder_text:
                raise RuntimeError(
                    "target_drift: bound completion span no longer matches placeholder identity"
                )
            replacement = normalized[target_id]
            replace_visible_span_across_runs(paragraph, start, end, replacement)
            expected_text = expected_text[:start] + replacement + expected_text[end:]
            target_verification.append(
                {
                    "target_id": target_id,
                    "location": location,
                    "placeholder_sha256": hashlib.sha256(
                        placeholder_text.encode("utf-8")
                    ).hexdigest(),
                    "replacement_sha256": hashlib.sha256(
                        replacement.encode("utf-8")
                    ).hexdigest(),
                }
            )
        expected_text_by_location[location] = expected_text
        if paragraph.text != expected_text:
            raise RuntimeError(
                "DOCX completion mutation did not produce the expected in-memory paragraph"
            )

    fd, raw_temp = tempfile.mkstemp(
        prefix=f".{destination.name}.",
        suffix=".tmp",
        dir=destination.parent,
    )
    os.close(fd)
    temp_path = Path(raw_temp)
    try:
        document.save(temp_path)
        temp_document = Document(temp_path)
        after_records = _records(temp_document)
        if len(after_records) != len(before_records):
            raise RuntimeError(
                "DOCX completion verification failed: paragraph topology changed"
            )
        changed_locations: list[str] = []
        for before, after in zip(before_records, after_records):
            if before["location"] != after["location"]:
                raise RuntimeError(
                    "DOCX completion verification failed: paragraph location order changed"
                )
            if len(before["run_text"]) != len(after["run_text"]):
                raise RuntimeError(
                    "DOCX completion verification failed: run topology changed"
                )
            if before["run_format"] != after["run_format"]:
                raise RuntimeError(
                    "DOCX completion verification failed: run formatting changed"
                )
            expected_text = expected_text_by_location.get(
                str(before["location"]),
                str(before["text"]),
            )
            if str(after["text"]) != expected_text:
                raise RuntimeError(
                    "DOCX completion verification failed: target or unrelated paragraph text changed unexpectedly"
                )
            if before["text"] != after["text"]:
                changed_locations.append(str(before["location"]))
        if set(changed_locations) != set(expected_text_by_location):
            raise RuntimeError(
                "DOCX completion verification failed: changed paragraph set differs from bound targets"
            )
        remaining = _completion_targets(after_records, enforce_max=False)
        if remaining:
            raise RuntimeError(
                "DOCX completion verification failed: placeholders remain after mutation"
            )

        source_after_temp = observe_file_identity(
            source,
            max_hash_bytes=MAX_FILE_BYTES,
        )
        if (
            compare_file_identities(current, source_after_temp).get("exact")
            is not True
        ):
            raise RuntimeError(
                "stale_source_evidence: DOCX source changed while completion copy was prepared"
            )
        if destination.exists():
            raise FileExistsError(
                "destination appeared during completion mutation; refusing overwrite"
            )
        _publish_no_overwrite(temp_path, destination)

        final_inspection = inspect_docx_completion_targets(destination)
        destination_identity = final_inspection.get("identity") or {}
        source_after = observe_file_identity(source, max_hash_bytes=MAX_FILE_BYTES)
        if compare_file_identities(current, source_after).get("exact") is not True:
            raise RuntimeError(
                "stale_source_evidence: DOCX source changed before final completion verification"
            )
        if final_inspection.get("ready") is not True:
            raise RuntimeError("DOCX completion final reopen/inspection failed")
        if int(final_inspection.get("target_count") or 0) != 0:
            raise RuntimeError(
                "DOCX completion final verification found remaining placeholders"
            )
        final_document = Document(destination)
        final_records = _records(final_document)
        if final_records != after_records:
            raise RuntimeError(
                "DOCX completion final reopen changed verified paragraph/run state"
            )
        expected_destination_fingerprint = _structure_fingerprint(after_records)
        if final_inspection.get(
            "structure_fingerprint"
        ) != expected_destination_fingerprint:
            raise RuntimeError(
                "DOCX completion final structure fingerprint changed after publication"
            )
        after_by_location = {
            str(record["location"]): record for record in final_records
        }
        for result in target_verification:
            location = str(result["location"])
            result["paragraph_sha256_after"] = hashlib.sha256(
                str(after_by_location[location]["text"]).encode("utf-8")
            ).hexdigest()

        return {
            "source_identity_before": current,
            "source_identity_after": source_after,
            "destination_identity": destination_identity,
            "source_unchanged": True,
            "destination_reopened": True,
            "source_structure_fingerprint": inspection["structure_fingerprint"],
            "expected_destination_structure_fingerprint": expected_destination_fingerprint,
            "target_ids": expected_ids,
            "target_count_before": len(expected_ids),
            "target_count_after": 0,
            "changed_locations": sorted(changed_locations),
            "run_topology_preserved": True,
            "format_preserved": True,
            "non_target_text_preserved": True,
            "all_placeholders_replaced": True,
            "target_verification": sorted(
                target_verification,
                key=lambda value: str(value["target_id"]),
            ),
        }
    except Exception:
        if temp_path.exists():
            temp_path.unlink(missing_ok=True)
        raise

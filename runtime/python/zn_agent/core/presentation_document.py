from __future__ import annotations

"""Deterministic bounded PPTX inspection and PresentationSpec v1 export."""

import hashlib
import json
import os
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Mapping, Sequence

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .file_identity import observe_file_identity
from .presentation_spec import (
    MAX_BULLETS as SPEC_MAX_BULLETS,
    MAX_BULLET_CHARS as SPEC_MAX_BULLET_CHARS,
    MAX_PRESENTATION_SLIDES,
    MAX_TITLE_CHARS as SPEC_MAX_TITLE_CHARS,
    normalize_presentation_spec,
    presentation_visible_texts,
)

MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_PACKAGE_PARTS = 4096
MAX_SLIDES = 40
MAX_TITLE_CHARS = max(240, SPEC_MAX_TITLE_CHARS)
MAX_SUBTITLE_CHARS = 600
MAX_BULLETS_PER_SLIDE = max(24, SPEC_MAX_BULLETS)
MAX_BULLET_CHARS = max(800, SPEC_MAX_BULLET_CHARS)
MAX_VISIBLE_CHARS = 64 * 1024

_UNSUPPORTED_PART_PREFIXES = (
    "ppt/activeX/",
    "ppt/embeddings/",
    "ppt/externalLinks/",
)
_UNSUPPORTED_PARTS = {"ppt/vbaProject.bin"}


def _block(identity: dict[str, Any], code: str, detail: str) -> dict[str, Any]:
    return {
        "ready": False,
        "blocker": code,
        "detail": str(detail)[:1000],
        "identity": identity,
    }


def _preflight(path: Path, identity: dict[str, Any]) -> tuple[bool, str | None]:
    if path.suffix.casefold() != ".pptx":
        return False, "only standard .pptx is supported"
    if (
        identity.get("exists") is not True
        or identity.get("type") != "file"
        or identity.get("stable") is not True
        or identity.get("digest_complete") is not True
    ):
        return False, "PPTX must be one stable regular file within the bounded size limit"
    try:
        with zipfile.ZipFile(path) as package:
            parts = package.infolist()
            names = {item.filename for item in parts}
            if len(parts) > MAX_PACKAGE_PARTS:
                return False, "PPTX package contains too many parts"
            expanded = sum(max(0, int(item.file_size)) for item in parts)
            if expanded > MAX_EXPANDED_BYTES:
                return False, "PPTX expanded package exceeds the bounded inspection limit"
            if "[Content_Types].xml" not in names or "ppt/presentation.xml" not in names:
                return False, "PPTX package is missing required presentation parts"
            for item in parts:
                name = item.filename.replace("\\", "/")
                if name.startswith("/") or name.startswith("../") or "/../" in name:
                    return False, "PPTX package contains an unsafe part path"
                if item.flag_bits & 0x1:
                    return False, "encrypted PPTX package parts are unsupported"
                if name in _UNSUPPORTED_PARTS or name.startswith(_UNSUPPORTED_PART_PREFIXES):
                    return False, f"unsupported PPTX package part: {name}"
    except (OSError, zipfile.BadZipFile) as exc:
        return False, f"invalid PPTX package: {type(exc).__name__}"
    return True, None


def _shape_texts(shape: Any) -> list[str]:
    texts: list[str] = []
    if bool(getattr(shape, "has_text_frame", False)):
        for paragraph in shape.text_frame.paragraphs:
            text = str(paragraph.text or "").strip()
            if text:
                texts.append(text)
    if bool(getattr(shape, "has_table", False)):
        for row in shape.table.rows:
            for cell in row.cells:
                text = str(cell.text or "").strip()
                if text:
                    texts.append(text)
    children = getattr(shape, "shapes", None)
    if children is not None:
        for child in children:
            texts.extend(_shape_texts(child))
    return texts


def _fingerprint(slides: list[dict[str, Any]]) -> str:
    payload = json.dumps(slides, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def inspect_pptx(path: str | Path) -> dict[str, Any]:
    source = Path(path).resolve(strict=False)
    identity = observe_file_identity(source, max_hash_bytes=MAX_FILE_BYTES)
    ok, detail = _preflight(source, identity)
    if not ok:
        return _block(identity, "unsupported_presentation_structure", detail or "unsupported PPTX")
    try:
        presentation = Presentation(source)
    except Exception as exc:
        return _block(identity, "pptx_open_failed", type(exc).__name__)
    if len(presentation.slides) > MAX_SLIDES:
        return _block(identity, "presentation_too_large", f"PPTX has more than {MAX_SLIDES} slides")

    slides: list[dict[str, Any]] = []
    visible_chars = 0
    for index, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            texts.extend(_shape_texts(shape))
        visible_chars += sum(len(text) for text in texts)
        if visible_chars > MAX_VISIBLE_CHARS:
            return _block(
                identity,
                "presentation_text_too_large",
                f"PPTX visible text exceeds {MAX_VISIBLE_CHARS} characters",
            )
        title_shape = slide.shapes.title
        title = str(title_shape.text or "").strip() if title_shape is not None else ""
        slides.append({"slide_index": index, "title": title, "texts": texts})

    return {
        "ready": True,
        "identity": identity,
        "slide_count": len(slides),
        "visible_text_chars": visible_chars,
        "slides": slides,
        "structure_fingerprint": _fingerprint(slides),
    }


def _bounded_text(value: Any, *, field_name: str, limit: int, required: bool) -> str:
    if not isinstance(value, str):
        raise ValueError(f"{field_name} must be a string")
    text = value.strip()
    if required and not text:
        raise ValueError(f"{field_name} must not be empty")
    if len(text) > limit:
        raise ValueError(f"{field_name} exceeds {limit} characters")
    return text


def _normalize_outline(
    *,
    title: str,
    subtitle: str,
    slides: Sequence[Mapping[str, Any]],
) -> tuple[str, str, list[dict[str, Any]]]:
    normalized_title = _bounded_text(
        title, field_name="presentation title", limit=MAX_TITLE_CHARS, required=True
    )
    normalized_subtitle = _bounded_text(
        subtitle, field_name="presentation subtitle", limit=MAX_SUBTITLE_CHARS, required=False
    )
    if isinstance(slides, (str, bytes, Mapping)) or not isinstance(slides, Sequence):
        raise ValueError("slides must be an ordered sequence of objects")
    if len(slides) > MAX_SLIDES - 1:
        raise ValueError(f"presentation supports at most {MAX_SLIDES - 1} content slides")

    normalized_slides: list[dict[str, Any]] = []
    total_chars = len(normalized_title) + len(normalized_subtitle)
    for index, raw in enumerate(slides, start=1):
        if not isinstance(raw, Mapping):
            raise ValueError(f"slide {index} must be an object")
        unexpected = set(raw) - {"title", "bullets"}
        if unexpected:
            raise ValueError(f"slide {index} contains unsupported fields: {sorted(unexpected)}")
        slide_title = _bounded_text(
            raw.get("title"), field_name=f"slide {index} title", limit=MAX_TITLE_CHARS, required=True
        )
        bullets_raw = raw.get("bullets", ())
        if isinstance(bullets_raw, (str, bytes, Mapping)) or not isinstance(bullets_raw, Sequence):
            raise ValueError(f"slide {index} bullets must be an ordered sequence of strings")
        if len(bullets_raw) > MAX_BULLETS_PER_SLIDE:
            raise ValueError(f"slide {index} supports at most {MAX_BULLETS_PER_SLIDE} bullets")
        bullets = [
            _bounded_text(
                item,
                field_name=f"slide {index} bullet {bullet_index}",
                limit=MAX_BULLET_CHARS,
                required=True,
            )
            for bullet_index, item in enumerate(bullets_raw, start=1)
        ]
        total_chars += len(slide_title) + sum(len(item) for item in bullets)
        if total_chars > MAX_VISIBLE_CHARS:
            raise ValueError(f"presentation outline exceeds {MAX_VISIBLE_CHARS} visible characters")
        normalized_slides.append({"title": slide_title, "bullets": bullets})
    return normalized_title, normalized_subtitle, normalized_slides


def _publish_no_overwrite(temp_path: Path, destination: Path) -> None:
    try:
        os.link(temp_path, destination)
    except FileExistsError:
        raise FileExistsError("output_collision: destination already exists; refusing overwrite")
    temp_path.unlink()


def _theme(theme: str) -> dict[str, str]:
    if theme == "light-clean":
        return {
            "bg": "F5F7FB",
            "panel": "FFFFFF",
            "text": "172033",
            "muted": "5B6577",
            "accent": "4C6FFF",
            "line": "D9E0EC",
        }
    return {
        "bg": "07111F",
        "panel": "0F2038",
        "text": "F5F8FF",
        "muted": "9EB2CB",
        "accent": "67E8F9",
        "line": "274564",
    }


def _fill_background(slide: Any, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(color)


def _add_text(
    slide: Any,
    text: str,
    *,
    x: float,
    y: float,
    w: float,
    h: float,
    size: int,
    color: str,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> Any:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    run = paragraph.runs[0]
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return box


def _add_title(slide: Any, title: str, colors: Mapping[str, str]) -> None:
    _add_text(
        slide,
        title,
        x=0.8,
        y=0.55,
        w=11.7,
        h=0.8,
        size=28,
        color=colors["text"],
        bold=True,
    )
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.38), Inches(1.25), Inches(0.06))
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGBColor.from_string(colors["accent"])
    accent.line.fill.background()


def _add_body_and_bullets(slide: Any, spec: Mapping[str, Any], colors: Mapping[str, str], *, width: float = 11.4) -> None:
    cursor = 1.75
    body = str(spec.get("body") or "").strip()
    if body:
        _add_text(
            slide,
            body,
            x=0.9,
            y=cursor,
            w=width,
            h=1.0,
            size=18,
            color=colors["muted"],
        )
        cursor += 1.05
    bullets = list(spec.get("bullets") or [])
    for bullet in bullets:
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(1.0),
            Inches(cursor + 0.22),
            Inches(0.09),
            Inches(0.09),
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = RGBColor.from_string(colors["accent"])
        dot.line.fill.background()
        _add_text(
            slide,
            str(bullet),
            x=1.22,
            y=cursor,
            w=width - 0.42,
            h=0.58,
            size=17,
            color=colors["text"],
        )
        cursor += 0.65


def _render_title_slide(slide: Any, spec: Mapping[str, Any], colors: Mapping[str, str]) -> None:
    _add_text(
        slide,
        str(spec["title"]),
        x=1.0,
        y=2.15,
        w=11.3,
        h=1.25,
        size=38,
        color=colors["text"],
        bold=True,
    )
    body = str(spec.get("body") or "").strip()
    if body:
        _add_text(
            slide,
            body,
            x=1.05,
            y=3.45,
            w=9.7,
            h=0.85,
            size=20,
            color=colors["muted"],
        )
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.05), Inches(4.55), Inches(2.0), Inches(0.08))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor.from_string(colors["accent"])
    line.line.fill.background()


def _render_architecture(slide: Any, spec: Mapping[str, Any], colors: Mapping[str, str]) -> None:
    _add_title(slide, str(spec["title"]), colors)
    diagram = dict(spec.get("diagram") or {})
    nodes = list(diagram.get("nodes") or [])
    edges = list(diagram.get("edges") or [])
    positions: dict[str, tuple[float, float, float, float]] = {}
    columns = 3 if len(nodes) > 4 else 2
    node_w, node_h = (3.2, 1.15) if columns == 3 else (4.2, 1.15)
    gap_x = 0.65
    start_x = 0.85 if columns == 3 else 1.55
    for index, node in enumerate(nodes):
        row, col = divmod(index, columns)
        x = start_x + col * (node_w + gap_x)
        y = 1.9 + row * 1.75
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(node_w),
            Inches(node_h),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(colors["panel"])
        shape.line.color.rgb = RGBColor.from_string(colors["line"])
        frame = shape.text_frame
        frame.clear()
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = frame.paragraphs[0]
        paragraph.text = str(node["label"])
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.runs[0]
        run.font.name = "Aptos"
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(colors["text"])
        positions[str(node["id"])] = (x, y, node_w, node_h)

    for edge in edges:
        source = positions[str(edge["from"])]
        target = positions[str(edge["to"])]
        x1, y1 = source[0] + source[2] / 2, source[1] + source[3] / 2
        x2, y2 = target[0] + target[2] / 2, target[1] + target[3] / 2
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(x1),
            Inches(y1),
            Inches(x2),
            Inches(y2),
        )
        connector.line.color.rgb = RGBColor.from_string(colors["accent"])
        connector.line.width = Pt(1.5)


def _render_image_right(
    slide: Any,
    spec: Mapping[str, Any],
    colors: Mapping[str, str],
    image_paths: Mapping[str, str],
) -> None:
    _add_title(slide, str(spec["title"]), colors)
    _add_body_and_bullets(slide, spec, colors, width=5.5)
    artifact_id = str(spec["image_artifact_id"])
    raw_path = str(image_paths.get(artifact_id) or "").strip()
    if not raw_path:
        raise ValueError(f"image artifact is not available for export: {artifact_id}")
    path = Path(raw_path).resolve(strict=True)
    if path.suffix.casefold() not in {".png", ".jpg", ".jpeg", ".webp"}:
        raise ValueError(f"unsupported presentation image type: {path.suffix}")
    slide.shapes.add_picture(str(path), Inches(7.1), Inches(1.75), width=Inches(5.25), height=Inches(4.75))


def _render_spec_presentation(
    spec: Mapping[str, Any],
    *,
    image_paths: Mapping[str, str],
) -> Presentation:
    normalized = normalize_presentation_spec(spec)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    colors = _theme(str(normalized["theme"]))
    for slide_spec in normalized["slides"]:
        slide = presentation.slides.add_slide(blank)
        _fill_background(slide, colors["bg"])
        layout = str(slide_spec["layout"])
        if layout == "title":
            _render_title_slide(slide, slide_spec, colors)
        elif layout == "architecture":
            _render_architecture(slide, slide_spec, colors)
        elif layout == "image-right":
            _render_image_right(slide, slide_spec, colors, image_paths)
        else:
            _add_title(slide, str(slide_spec["title"]), colors)
            _add_body_and_bullets(slide, slide_spec, colors)
    return presentation


def create_pptx_from_spec(
    destination_path: str | Path,
    *,
    spec: Mapping[str, Any],
    image_paths: Mapping[str, str] | None = None,
) -> dict[str, Any]:
    destination = Path(destination_path).resolve(strict=False)
    if destination.suffix.casefold() != ".pptx":
        raise ValueError("PPTX destination must use the .pptx extension")
    if not destination.parent.exists() or not destination.parent.is_dir():
        raise FileNotFoundError("PPTX destination parent directory does not exist")
    if destination.exists():
        raise FileExistsError("output_collision: destination already exists; refusing overwrite")
    normalized = normalize_presentation_spec(spec)
    expected_texts = presentation_visible_texts(normalized)
    if len(expected_texts) > MAX_PRESENTATION_SLIDES:
        raise ValueError("PresentationSpec exceeds bounded exporter slide count")

    fd, raw_temp = tempfile.mkstemp(
        prefix=f".{destination.name}.",
        suffix=".pptx",
        dir=destination.parent,
    )
    os.close(fd)
    temp_path = Path(raw_temp)
    try:
        presentation = _render_spec_presentation(
            normalized,
            image_paths=dict(image_paths or {}),
        )
        presentation.save(temp_path)
        temp_inspection = inspect_pptx(temp_path)
        if temp_inspection.get("ready") is not True:
            raise RuntimeError(
                f"PPTX generated package failed reopen: {temp_inspection.get('detail')}"
            )
        observed_texts = [list(item.get("texts") or ()) for item in temp_inspection["slides"]]
        if observed_texts != expected_texts:
            raise RuntimeError("PPTX generated visible text differs from PresentationSpec")
        if int(temp_inspection["slide_count"]) != len(expected_texts):
            raise RuntimeError("PPTX generated slide count differs from PresentationSpec")
        if destination.exists():
            raise FileExistsError("output_collision: destination appeared during PPTX generation")
        _publish_no_overwrite(temp_path, destination)

        final_inspection = inspect_pptx(destination)
        if final_inspection.get("ready") is not True:
            raise RuntimeError("PPTX final fresh reopen/inspection failed")
        final_texts = [list(item.get("texts") or ()) for item in final_inspection["slides"]]
        if final_texts != expected_texts:
            raise RuntimeError("PPTX final visible text differs from PresentationSpec")
        if final_inspection.get("structure_fingerprint") != temp_inspection.get("structure_fingerprint"):
            raise RuntimeError("PPTX final structure fingerprint differs from verified temp output")
        return {
            "verified": True,
            "destination": str(destination),
            "identity": final_inspection["identity"],
            "slide_count": final_inspection["slide_count"],
            "visible_text_chars": final_inspection["visible_text_chars"],
            "structure_fingerprint": final_inspection["structure_fingerprint"],
        }
    except Exception:
        if temp_path.exists():
            temp_path.unlink(missing_ok=True)
        raise


def create_pptx_from_outline(
    destination_path: str | Path,
    *,
    title: str,
    subtitle: str = "",
    slides: Sequence[Mapping[str, Any]] = (),
) -> dict[str, Any]:
    destination = Path(destination_path).resolve(strict=False)
    if destination.suffix.casefold() != ".pptx":
        raise ValueError("PPTX destination must use the .pptx extension")
    if not destination.parent.exists() or not destination.parent.is_dir():
        raise FileNotFoundError("PPTX destination parent directory does not exist")
    if destination.exists():
        raise FileExistsError("output_collision: destination already exists; refusing overwrite")

    normalized_title, normalized_subtitle, normalized_slides = _normalize_outline(
        title=title,
        subtitle=subtitle,
        slides=slides,
    )
    expected_texts = [[text for text in (normalized_title, normalized_subtitle) if text]]
    expected_texts.extend(
        [str(slide["title"]), *[str(item) for item in slide["bullets"]]]
        for slide in normalized_slides
    )

    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    if title_slide.shapes.title is None:
        raise RuntimeError("default PPTX title layout has no title placeholder")
    title_slide.shapes.title.text = normalized_title
    if normalized_subtitle:
        if len(title_slide.placeholders) < 2:
            raise RuntimeError("default PPTX title layout has no subtitle placeholder")
        title_slide.placeholders[1].text = normalized_subtitle

    for spec in normalized_slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        if slide.shapes.title is None or len(slide.placeholders) < 2:
            raise RuntimeError("default PPTX content layout lacks required placeholders")
        slide.shapes.title.text = str(spec["title"])
        text_frame = slide.placeholders[1].text_frame
        text_frame.clear()
        bullets = list(spec["bullets"])
        if bullets:
            text_frame.paragraphs[0].text = bullets[0]
            for bullet in bullets[1:]:
                paragraph = text_frame.add_paragraph()
                paragraph.text = bullet

    fd, raw_temp = tempfile.mkstemp(
        prefix=f".{destination.name}.",
        suffix=".pptx",
        dir=destination.parent,
    )
    os.close(fd)
    temp_path = Path(raw_temp)
    try:
        presentation.save(temp_path)
        temp_inspection = inspect_pptx(temp_path)
        if temp_inspection.get("ready") is not True:
            raise RuntimeError(
                f"PPTX generated package failed reopen: {temp_inspection.get('detail')}"
            )
        observed_texts = [list(item.get("texts") or ()) for item in temp_inspection["slides"]]
        if observed_texts != expected_texts:
            raise RuntimeError("PPTX generated visible text differs from the admitted outline")
        if int(temp_inspection["slide_count"]) != len(expected_texts):
            raise RuntimeError("PPTX generated slide count differs from the admitted outline")
        if destination.exists():
            raise FileExistsError("output_collision: destination appeared during PPTX generation")
        _publish_no_overwrite(temp_path, destination)

        final_inspection = inspect_pptx(destination)
        if final_inspection.get("ready") is not True:
            raise RuntimeError("PPTX final fresh reopen/inspection failed")
        final_texts = [list(item.get("texts") or ()) for item in final_inspection["slides"]]
        if final_texts != expected_texts:
            raise RuntimeError("PPTX final visible text differs from the admitted outline")
        if final_inspection.get("structure_fingerprint") != temp_inspection.get(
            "structure_fingerprint"
        ):
            raise RuntimeError("PPTX final structure fingerprint differs from verified temp output")
        return {
            "verified": True,
            "destination": str(destination),
            "identity": final_inspection["identity"],
            "slide_count": final_inspection["slide_count"],
            "visible_text_chars": final_inspection["visible_text_chars"],
            "structure_fingerprint": final_inspection["structure_fingerprint"],
        }
    except Exception:
        if temp_path.exists():
            temp_path.unlink(missing_ok=True)
        raise

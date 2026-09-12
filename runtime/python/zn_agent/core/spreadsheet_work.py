from __future__ import annotations

"""Deterministic bounded XLSX inspection, cleanup-copy, and append-copy mutation."""

import hashlib
import json
import math
import os
import tempfile
import zipfile
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Sequence
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.cell.rich_text import CellRichText

from .file_identity import compare_file_identities, observe_file_identity

MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_ROWS = 2_000
MAX_COLUMNS = 64
AMOUNT_NUMBER_FORMAT = "#,##0.00"

# E2E-11 deliberately inherits the Browser structured-table bounds. The table
# row cap includes its header, so append data itself is at most 63 rows.
MAX_BROWSER_TABLE_ROWS = 64
MAX_BROWSER_TABLE_COLUMNS = 32
MAX_BROWSER_CELL_TEXT = 512

_UNSUPPORTED_PART_PREFIXES = (
    "xl/charts/",
    "xl/drawings/",
    "xl/embeddings/",
    "xl/externalLinks/",
    "xl/media/",
    "xl/pivotCache/",
    "xl/pivotTables/",
    "xl/queryTables/",
    "xl/slicers/",
    "xl/tables/",
)
_UNSUPPORTED_PARTS = {"xl/vbaProject.bin", "xl/connections.xml"}
_WORKBOOK_UNSUPPORTED = {"definedName", "externalReference"}
_SHEET_UNSUPPORTED = {
    "autoFilter",
    "conditionalFormatting",
    "controls",
    "dataValidations",
    "drawing",
    "f",
    "hyperlink",
    "legacyDrawing",
    "mergeCell",
    "oleObject",
    "sheetProtection",
    "tablePart",
}


def _localname(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag.rsplit(":", 1)[-1]


def _block(identity: dict[str, Any], code: str, detail: str) -> dict[str, Any]:
    return {"ready": False, "blocker": code, "detail": str(detail)[:1000], "identity": identity}


def _preflight(path: Path, identity: dict[str, Any]) -> tuple[bool, str | None]:
    if path.suffix.casefold() != ".xlsx":
        return False, "only standard .xlsx is supported"
    if (
        identity.get("exists") is not True
        or identity.get("type") != "file"
        or identity.get("stable") is not True
        or identity.get("digest_complete") is not True
    ):
        return False, "source file identity is not a stable complete regular-file observation"
    if int(identity.get("size_bytes") or 0) > MAX_FILE_BYTES:
        return False, "XLSX exceeds the bounded file-size limit"
    if not zipfile.is_zipfile(path):
        return False, "encrypted, legacy, corrupt, or non-OOXML spreadsheet packages are unsupported"
    try:
        with zipfile.ZipFile(path, "r") as package:
            names = set(package.namelist())
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                return False, "required SpreadsheetML package parts are missing"
            for name in names:
                if name in _UNSUPPORTED_PARTS or any(
                    name.startswith(prefix) for prefix in _UNSUPPORTED_PART_PREFIXES
                ):
                    return False, f"unsupported XLSX package part: {name}"
            workbook_root = ET.fromstring(package.read("xl/workbook.xml"))
            bad_workbook = sorted(
                {_localname(el.tag) for el in workbook_root.iter()} & _WORKBOOK_UNSUPPORTED
            )
            if bad_workbook:
                return False, "unsupported workbook structures: " + ", ".join(bad_workbook)
            for element in workbook_root.iter():
                if _localname(element.tag) == "workbookProtection" and element.attrib:
                    return False, "protected workbooks are outside XLSX v1 scope"
            sheet_names = sorted(
                name
                for name in names
                if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
            )
            if len(sheet_names) != 1:
                return False, "XLSX v1 requires exactly one worksheet"
            sheet_root = ET.fromstring(package.read(sheet_names[0]))
            bad_sheet = sorted(
                {_localname(el.tag) for el in sheet_root.iter()} & _SHEET_UNSUPPORTED
            )
            if bad_sheet:
                return False, "unsupported worksheet structures: " + ", ".join(bad_sheet)
    except (OSError, KeyError, zipfile.BadZipFile, ET.ParseError) as exc:
        return False, f"XLSX package preflight failed: {type(exc).__name__}: {exc}"
    return True, None


def _semantic_value(value: Any) -> tuple[str, str]:
    if value is None:
        return ("none", "")
    if isinstance(value, bool):
        return ("bool", "1" if value else "0")
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError("non-finite numeric values are outside XLSX v1 scope")
        try:
            number = Decimal(str(value))
        except InvalidOperation as exc:
            raise ValueError("invalid numeric cell value") from exc
        if not number.is_finite():
            raise ValueError("non-finite numeric values are outside XLSX v1 scope")
        if number == 0:
            number = Decimal(0)
        return ("number", format(number.normalize(), "f"))
    if isinstance(value, str):
        return ("text", value)
    if isinstance(value, datetime):
        return ("datetime", value.isoformat())
    if isinstance(value, date):
        return ("date", value.isoformat())
    if isinstance(value, time):
        return ("time", value.isoformat())
    raise ValueError(f"unsupported scalar cell type: {type(value).__name__}")


def _row_key(values: list[Any]) -> tuple[tuple[str, str], ...]:
    return tuple(_semantic_value(value) for value in values)


def _resolve_amount_header(headers: list[str]) -> tuple[int | None, str | None]:
    normalized = [value.strip().casefold() for value in headers]
    exact = [index for index, value in enumerate(normalized) if value in {"金额", "amount"}]
    if len(exact) == 1:
        return exact[0] + 1, None
    amount_like = [
        index
        for index, value in enumerate(normalized)
        if "金额" in value or "amount" in value
    ]
    if len(exact) > 1 or amount_like:
        return None, "amount column is ambiguous; user confirmation is required"
    return None, "no unique explicit 金额/Amount column exists"


def _load_supported(path: Path, *, rich_text: bool = False):
    return load_workbook(
        path,
        read_only=False,
        data_only=False,
        keep_links=False,
        rich_text=rich_text,
    )


def _snapshot(path: Path) -> dict[str, Any]:
    """E2E-10 cleanup-specific snapshot. Keep public behavior unchanged."""
    workbook = _load_supported(path)
    if len(workbook.worksheets) != 1:
        raise ValueError("XLSX v1 requires exactly one worksheet")
    sheet = workbook.worksheets[0]
    if sheet.max_row < 2 or sheet.max_column < 1:
        raise ValueError("worksheet must contain one header row and at least one data row")
    if sheet.max_row > MAX_ROWS + 1 or sheet.max_column > MAX_COLUMNS:
        raise ValueError("worksheet used range exceeds bounded XLSX v1 limits")
    if sheet.merged_cells.ranges:
        raise ValueError("merged cells are outside XLSX v1 scope")
    if any(bool(dim.hidden) for dim in sheet.row_dimensions.values()) or any(
        bool(dim.hidden) for dim in sheet.column_dimensions.values()
    ):
        raise ValueError("hidden rows or columns are outside XLSX v1 scope")

    headers: list[str] = []
    for col in range(1, sheet.max_column + 1):
        value = sheet.cell(row=1, column=col).value
        if not isinstance(value, str) or not value.strip():
            raise ValueError("header row must contain non-empty text in every used column")
        headers.append(value)
    if len({header.strip().casefold() for header in headers}) != len(headers):
        raise ValueError("duplicate header names are outside XLSX v1 scope")
    amount_col, amount_error = _resolve_amount_header(headers)
    if amount_col is None:
        raise LookupError(amount_error or "amount column is unavailable")

    rows: list[list[Any]] = []
    row_keys: list[tuple[tuple[str, str], ...]] = []
    amount_formats: list[str] = []
    for row_index in range(2, sheet.max_row + 1):
        values = [
            sheet.cell(row=row_index, column=col).value
            for col in range(1, sheet.max_column + 1)
        ]
        if all(value is None for value in values):
            raise ValueError("blank rows inside the used data region are outside XLSX v1 scope")
        if any(value is None for value in values):
            raise ValueError("sparse rows inside the used data region are outside XLSX v1 scope")
        for col, value in enumerate(values, start=1):
            cell: Cell = sheet.cell(row=row_index, column=col)
            if cell.data_type == "f":
                raise ValueError("formula cells are outside XLSX v1 scope")
            _semantic_value(value)
        amount = sheet.cell(row=row_index, column=amount_col)
        if amount.data_type != "n" or isinstance(amount.value, bool) or not isinstance(
            amount.value, (int, float, Decimal)
        ):
            raise TypeError("amount cells must be numeric cells")
        _semantic_value(amount.value)
        rows.append(values)
        row_keys.append(_row_key(values))
        amount_formats.append(str(amount.number_format))

    seen: set[tuple[tuple[str, str], ...]] = set()
    duplicate_indices: list[int] = []
    retained_indices: list[int] = []
    for index, key in enumerate(row_keys):
        if key in seen:
            duplicate_indices.append(index)
        else:
            seen.add(key)
            retained_indices.append(index)

    retained_rows = [rows[index] for index in retained_indices]
    canonical = {
        "sheet_name": sheet.title,
        "headers": headers,
        "amount_col": amount_col,
        "rows": [
            [list(_semantic_value(value)) for value in row] for row in retained_rows
        ],
    }
    non_target = {
        "headers": headers,
        "rows": [
            [
                list(_semantic_value(value))
                for col, value in enumerate(row, start=1)
                if col != amount_col
            ]
            for row in retained_rows
        ],
    }
    amount_values = [
        list(_semantic_value(row[amount_col - 1])) for row in retained_rows
    ]
    return {
        "workbook": workbook,
        "sheet": sheet,
        "headers": headers,
        "amount_col": amount_col,
        "rows": rows,
        "row_keys": row_keys,
        "retained_indices": retained_indices,
        "duplicate_indices": duplicate_indices,
        "retained_rows": retained_rows,
        "amount_formats": amount_formats,
        "data_fingerprint": hashlib.sha256(
            json.dumps(
                canonical,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ).encode("utf-8")
        ).hexdigest(),
        "non_target_fingerprint": hashlib.sha256(
            json.dumps(
                non_target,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ).encode("utf-8")
        ).hexdigest(),
        "amount_semantic_fingerprint": hashlib.sha256(
            json.dumps(
                amount_values,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ).encode("utf-8")
        ).hexdigest(),
    }


def inspect_xlsx(path: str | Path) -> dict[str, Any]:
    source = Path(path).resolve(strict=False)
    identity = observe_file_identity(source, max_hash_bytes=MAX_FILE_BYTES)
    ok, detail = _preflight(source, identity)
    if not ok:
        return _block(
            identity,
            "unsupported_workbook_structure",
            detail or "unsupported XLSX",
        )
    try:
        snap = _snapshot(source)
    except LookupError as exc:
        return _block(identity, "ambiguous_amount_column", str(exc))
    except (TypeError, ValueError) as exc:
        detail = str(exc)
        code = "unsupported_workbook_structure" if any(
            token in detail.lower()
            for token in ("formula", "merged", "hidden", "scope", "sparse", "blank")
        ) else "invalid_spreadsheet_data"
        return _block(identity, code, detail)
    formats = list(snap["amount_formats"])
    return {
        "ready": True,
        "blocker": None,
        "identity": identity,
        "worksheet_count": 1,
        "sheet_name": snap["sheet"].title,
        "headers": list(snap["headers"]),
        "data_row_count": len(snap["rows"]),
        "duplicate_row_count": len(snap["duplicate_indices"]),
        "retained_row_count": len(snap["retained_rows"]),
        "amount_column_index": int(snap["amount_col"]),
        "amount_header": snap["headers"][snap["amount_col"] - 1],
        "amount_number_formats": sorted(set(formats)),
        "amount_format_uniform": len(set(formats)) == 1,
        "data_fingerprint": snap["data_fingerprint"],
        "non_target_fingerprint": snap["non_target_fingerprint"],
        "amount_semantic_fingerprint": snap["amount_semantic_fingerprint"],
    }


def _append_cell_evidence(cell: Cell) -> list[Any]:
    if isinstance(cell.value, CellRichText):
        raise ValueError("rich text cells are outside XLSX append-copy scope")
    if cell.data_type == "f":
        raise ValueError("formula cells are outside XLSX append-copy scope")
    return [
        list(_semantic_value(cell.value)),
        str(cell.data_type or ""),
        str(cell.number_format or ""),
    ]


def _append_region_fingerprint(sheet: Any, *, max_row: int, max_column: int) -> str:
    canonical = {
        "sheet_name": sheet.title,
        "max_row": int(max_row),
        "max_column": int(max_column),
        "cells": [
            [
                _append_cell_evidence(sheet.cell(row=row, column=column))
                for column in range(1, max_column + 1)
            ]
            for row in range(1, max_row + 1)
        ],
    }
    return hashlib.sha256(
        json.dumps(
            canonical,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()


def _append_snapshot(path: Path) -> dict[str, Any]:
    """Snapshot the conservative one-sheet append target without E2E-10 semantics."""
    workbook = _load_supported(path, rich_text=True)
    if len(workbook.worksheets) != 1:
        raise ValueError("XLSX append-copy requires exactly one worksheet")
    sheet = workbook.worksheets[0]
    if sheet.max_row < 1 or sheet.max_column < 1:
        raise ValueError("worksheet must contain one header row")
    if sheet.max_row > MAX_ROWS + 1 or sheet.max_column > MAX_COLUMNS:
        raise ValueError("worksheet used range exceeds bounded XLSX v1 limits")
    if sheet.merged_cells.ranges:
        raise ValueError("merged cells are outside XLSX append-copy scope")
    if any(bool(dim.hidden) for dim in sheet.row_dimensions.values()) or any(
        bool(dim.hidden) for dim in sheet.column_dimensions.values()
    ):
        raise ValueError("hidden rows or columns are outside XLSX append-copy scope")

    headers: list[str] = []
    for column in range(1, sheet.max_column + 1):
        cell = sheet.cell(row=1, column=column)
        value = cell.value
        if isinstance(value, CellRichText):
            raise ValueError("rich text cells are outside XLSX append-copy scope")
        if cell.data_type == "f" or not isinstance(value, str) or not value.strip():
            raise ValueError(
                "header row must contain non-empty stable strings in every used column"
            )
        headers.append(value)
    if len(set(headers)) != len(headers) or len({value.strip() for value in headers}) != len(headers):
        raise ValueError("header row must contain unique stable strings")

    rows: list[list[Any]] = []
    for row_index in range(2, sheet.max_row + 1):
        values = [
            sheet.cell(row=row_index, column=column).value
            for column in range(1, sheet.max_column + 1)
        ]
        if any(value is None for value in values):
            raise ValueError(
                "blank or sparse rows inside the used data region are outside XLSX append-copy scope"
            )
        for column in range(1, sheet.max_column + 1):
            _append_cell_evidence(sheet.cell(row=row_index, column=column))
        rows.append(values)

    fingerprint = _append_region_fingerprint(
        sheet,
        max_row=sheet.max_row,
        max_column=sheet.max_column,
    )
    return {
        "workbook": workbook,
        "sheet": sheet,
        "headers": headers,
        "rows": rows,
        "data_row_count": len(rows),
        "semantic_fingerprint": fingerprint,
    }


def inspect_xlsx_append_target(path: str | Path) -> dict[str, Any]:
    """Inspect one exact XLSX as the bounded append target for Browser data."""
    source = Path(path).resolve(strict=False)
    identity = observe_file_identity(source, max_hash_bytes=MAX_FILE_BYTES)
    ok, detail = _preflight(source, identity)
    if not ok:
        return _block(
            identity,
            "unsupported_workbook_structure",
            detail or "unsupported XLSX",
        )
    try:
        snap = _append_snapshot(source)
    except (TypeError, ValueError) as exc:
        return _block(identity, "unsupported_workbook_structure", str(exc))
    return {
        "ready": True,
        "blocker": None,
        "identity": identity,
        "worksheet_count": 1,
        "sheet_name": snap["sheet"].title,
        "headers": list(snap["headers"]),
        "existing_row_count": int(snap["data_row_count"]),
        "semantic_fingerprint": snap["semantic_fingerprint"],
        "pre_existing_semantic_fingerprint": snap["semantic_fingerprint"],
    }


def _publish_no_overwrite(temp_path: Path, destination: Path) -> None:
    try:
        os.link(temp_path, destination)
    except FileExistsError:
        raise FileExistsError("destination already exists; refusing overwrite")
    temp_path.unlink()


def _browser_append_input(
    headers: Sequence[str],
    rows: Iterable[Sequence[str]],
) -> tuple[list[str], list[list[str]]]:
    if isinstance(headers, (str, bytes)) or not isinstance(headers, Sequence):
        raise ValueError("headers must be an ordered sequence of strings")
    normalized_headers = list(headers)
    if not normalized_headers or len(normalized_headers) > MAX_BROWSER_TABLE_COLUMNS:
        raise ValueError("headers exceed the Browser table column bound")
    if any(not isinstance(value, str) or not value for value in normalized_headers):
        raise ValueError("headers must contain non-empty strings")
    if len(set(normalized_headers)) != len(normalized_headers):
        raise ValueError("headers must be unique")

    if isinstance(rows, (str, bytes)):
        raise ValueError("rows must be an ordered sequence of rows")
    normalized_rows: list[list[str]] = []
    for raw_row in rows:
        if isinstance(raw_row, (str, bytes)) or not isinstance(raw_row, Sequence):
            raise ValueError("each appended row must be an ordered sequence")
        row = list(raw_row)
        if len(row) != len(normalized_headers):
            raise ValueError("appended rows must be rectangular and match the header width")
        if any(not isinstance(value, str) for value in row):
            raise ValueError("Browser spreadsheet cells must be strings")
        if any(len(value) > MAX_BROWSER_CELL_TEXT for value in row):
            raise ValueError("Browser spreadsheet cell text exceeds the bounded limit")
        normalized_rows.append(row)
    if not normalized_rows:
        raise ValueError("at least one Browser data row is required")
    if len(normalized_rows) + 1 > MAX_BROWSER_TABLE_ROWS:
        raise ValueError("Browser spreadsheet rows exceed the bounded table row limit")
    return normalized_headers, normalized_rows


def append_xlsx_rows_copy(
    source_path: str | Path,
    destination_path: str | Path,
    *,
    precondition_identity: dict[str, Any],
    headers: Sequence[str],
    rows: Iterable[Sequence[str]],
) -> dict[str, Any]:
    """Append Browser strings to a verified copy, preserving the exact source."""
    source = Path(source_path).resolve(strict=True)
    destination = Path(destination_path).resolve(strict=False)
    if source == destination:
        raise ValueError("XLSX append-copy destination must differ from source")
    if destination.exists():
        raise FileExistsError("output_collision: destination already exists; refusing overwrite")
    if not destination.parent.is_dir():
        raise ValueError("destination parent is unavailable")

    requested_headers, requested_rows = _browser_append_input(headers, rows)
    current = observe_file_identity(source, max_hash_bytes=MAX_FILE_BYTES)
    if compare_file_identities(precondition_identity, current).get("exact") is not True:
        raise RuntimeError("stale_source_evidence: exact XLSX identity changed before mutation")
    inspection = inspect_xlsx_append_target(source)
    if inspection.get("ready") is not True:
        raise RuntimeError(f"{inspection.get('blocker')}: {inspection.get('detail')}")
    if list(inspection.get("headers") or []) != requested_headers:
        raise ValueError("spreadsheet_header_mismatch: Browser and XLSX headers differ")

    before = _append_snapshot(source)
    before_row_count = int(before["data_row_count"])
    if before_row_count + len(requested_rows) > MAX_ROWS:
        raise ValueError("final XLSX row count would exceed the bounded XLSX v1 limit")
    workbook = before["workbook"]
    sheet = before["sheet"]
    original_max_row = sheet.max_row
    original_max_column = sheet.max_column
    existing_fingerprint = before["semantic_fingerprint"]

    for row in requested_rows:
        sheet.append(row)
        # openpyxl interprets leading '=' as a formula on append. Browser table
        # cells are literal observed strings, so force the cell data type back to
        # string before save rather than inventing spreadsheet semantics.
        row_index = sheet.max_row
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=column)
            cell.value = value
            cell.data_type = "s"

    fd, raw_temp = tempfile.mkstemp(
        prefix=f".{destination.name}.",
        suffix=".xlsx",
        dir=destination.parent,
    )
    os.close(fd)
    temp_path = Path(raw_temp)
    try:
        workbook.save(temp_path)
        temp_identity = observe_file_identity(temp_path, max_hash_bytes=MAX_FILE_BYTES)
        ok, detail = _preflight(temp_path, temp_identity)
        if not ok:
            raise RuntimeError(f"XLSX append-copy reopen preflight failed: {detail}")
        after = _append_snapshot(temp_path)
        if after["sheet"].title != before["sheet"].title:
            raise RuntimeError("XLSX append-copy verification failed: sheet name changed")
        if after["headers"] != before["headers"]:
            raise RuntimeError("XLSX append-copy verification failed: header changed")
        if int(after["data_row_count"]) != before_row_count + len(requested_rows):
            raise RuntimeError("XLSX append-copy verification failed: final row count is wrong")
        after_existing = _append_region_fingerprint(
            after["sheet"],
            max_row=original_max_row,
            max_column=original_max_column,
        )
        if after_existing != existing_fingerprint:
            raise RuntimeError(
                "XLSX append-copy verification failed: existing content or formatting changed"
            )
        appended_values: list[list[str]] = []
        for row_index in range(original_max_row + 1, after["sheet"].max_row + 1):
            values: list[str] = []
            for column in range(1, original_max_column + 1):
                cell = after["sheet"].cell(row=row_index, column=column)
                if cell.data_type != "s" or not isinstance(cell.value, str):
                    raise RuntimeError(
                        "XLSX append-copy verification failed: imported Browser cell is not a string"
                    )
                values.append(cell.value)
            appended_values.append(values)
        if appended_values != requested_rows:
            raise RuntimeError(
                "XLSX append-copy verification failed: appended rows differ from Browser evidence"
            )

        source_after_temp = observe_file_identity(source, max_hash_bytes=MAX_FILE_BYTES)
        if compare_file_identities(current, source_after_temp).get("exact") is not True:
            raise RuntimeError("stale_source_evidence: XLSX source changed while copy was prepared")
        if destination.exists():
            raise FileExistsError("output_collision: destination appeared during append-copy")
        _publish_no_overwrite(temp_path, destination)

        final_inspection = inspect_xlsx_append_target(destination)
        source_after = observe_file_identity(source, max_hash_bytes=MAX_FILE_BYTES)
        if compare_file_identities(current, source_after).get("exact") is not True:
            raise RuntimeError("stale_source_evidence: XLSX source changed before final verification")
        if final_inspection.get("ready") is not True:
            raise RuntimeError("XLSX append-copy final fresh reopen/inspection failed")
        if final_inspection.get("sheet_name") != before["sheet"].title:
            raise RuntimeError("XLSX append-copy final verification found a changed sheet name")
        if list(final_inspection.get("headers") or []) != requested_headers:
            raise RuntimeError("XLSX append-copy final verification found changed headers")
        if int(final_inspection.get("existing_row_count") or -1) != before_row_count + len(requested_rows):
            raise RuntimeError("XLSX append-copy final verification found an unexpected row count")
        if final_inspection.get("semantic_fingerprint") != after["semantic_fingerprint"]:
            raise RuntimeError("XLSX append-copy final semantic fingerprint changed after publication")

        return {
            "source_identity_before": current,
            "source_identity_after": source_after,
            "destination_identity": final_inspection["identity"],
            "source_unchanged": True,
            "destination_reopened": True,
            "sheet_name": before["sheet"].title,
            "headers": requested_headers,
            "original_data_rows": before_row_count,
            "imported_row_count": len(requested_rows),
            "final_data_rows": before_row_count + len(requested_rows),
            "existing_content_unchanged": True,
            "existing_semantic_fingerprint_before": existing_fingerprint,
            "existing_semantic_fingerprint_after": after_existing,
            "appended_rows_exact": True,
            "appended_rows_fingerprint": hashlib.sha256(
                json.dumps(
                    requested_rows,
                    ensure_ascii=False,
                    separators=(",", ":"),
                ).encode("utf-8")
            ).hexdigest(),
            "destination_semantic_fingerprint": after["semantic_fingerprint"],
        }
    except Exception:
        if temp_path.exists():
            temp_path.unlink(missing_ok=True)
        raise


def write_xlsx_copy(
    source_path: str | Path,
    destination_path: str | Path,
    *,
    precondition_identity: dict[str, Any],
    number_format: str = AMOUNT_NUMBER_FORMAT,
) -> dict[str, Any]:
    source = Path(source_path).resolve(strict=True)
    destination = Path(destination_path).resolve(strict=False)
    if source == destination:
        raise ValueError("XLSX copy destination must differ from source")
    if destination.exists():
        raise FileExistsError("output_collision: destination already exists; refusing overwrite")
    if not destination.parent.is_dir():
        raise ValueError("destination parent is unavailable")

    current = observe_file_identity(source, max_hash_bytes=MAX_FILE_BYTES)
    if compare_file_identities(precondition_identity, current).get("exact") is not True:
        raise RuntimeError("stale_source_evidence: exact XLSX identity changed before mutation")
    inspection = inspect_xlsx(source)
    if inspection.get("ready") is not True:
        raise RuntimeError(f"{inspection.get('blocker')}: {inspection.get('detail')}")
    before = _snapshot(source)
    workbook = before["workbook"]
    sheet = before["sheet"]
    amount_col = int(before["amount_col"])

    delete_rows = [index + 2 for index in before["duplicate_indices"]]
    for row_index in sorted(delete_rows, reverse=True):
        sheet.delete_rows(row_index, 1)
    for row_index in range(2, sheet.max_row + 1):
        amount = sheet.cell(row=row_index, column=amount_col)
        if amount.data_type != "n" or isinstance(amount.value, bool) or not isinstance(
            amount.value, (int, float, Decimal)
        ):
            raise RuntimeError("amount cell stopped being numeric during cleanup")
        amount.number_format = number_format

    fd, raw_temp = tempfile.mkstemp(
        prefix=f".{destination.name}.", suffix=".xlsx", dir=destination.parent
    )
    os.close(fd)
    temp_path = Path(raw_temp)
    try:
        workbook.save(temp_path)
        temp_identity = observe_file_identity(temp_path, max_hash_bytes=MAX_FILE_BYTES)
        ok, detail = _preflight(temp_path, temp_identity)
        if not ok:
            raise RuntimeError(f"XLSX reopen preflight failed: {detail}")
        after = _snapshot(temp_path)
        if after["headers"] != before["headers"]:
            raise RuntimeError("XLSX verification failed: header changed")
        if after["sheet"].title != before["sheet"].title:
            raise RuntimeError("XLSX verification failed: sheet name changed")
        if len(after["rows"]) != len(before["retained_rows"]):
            raise RuntimeError("XLSX verification failed: retained row count is wrong")
        if after["duplicate_indices"]:
            raise RuntimeError("XLSX verification failed: exact duplicate rows remain")
        if after["data_fingerprint"] != before["data_fingerprint"]:
            raise RuntimeError(
                "XLSX verification failed: retained business values or order changed"
            )
        if after["non_target_fingerprint"] != before["non_target_fingerprint"]:
            raise RuntimeError("XLSX verification failed: non-target column values changed")
        if after["amount_semantic_fingerprint"] != before["amount_semantic_fingerprint"]:
            raise RuntimeError("XLSX verification failed: amount numeric semantics changed")
        for row_index in range(2, after["sheet"].max_row + 1):
            amount = after["sheet"].cell(row=row_index, column=amount_col)
            if amount.data_type != "n" or isinstance(amount.value, bool) or not isinstance(
                amount.value, (int, float, Decimal)
            ):
                raise RuntimeError("XLSX verification failed: amount cell is no longer numeric")
            if amount.number_format != number_format:
                raise RuntimeError(
                    "XLSX verification failed: amount number_format is not uniform"
                )

        source_after_temp = observe_file_identity(source, max_hash_bytes=MAX_FILE_BYTES)
        if compare_file_identities(current, source_after_temp).get("exact") is not True:
            raise RuntimeError("XLSX source changed while cleaned copy was being prepared")
        if destination.exists():
            raise FileExistsError("output_collision: destination appeared during cleanup")
        _publish_no_overwrite(temp_path, destination)

        final_inspection = inspect_xlsx(destination)
        source_after = observe_file_identity(source, max_hash_bytes=MAX_FILE_BYTES)
        if compare_file_identities(current, source_after).get("exact") is not True:
            raise RuntimeError("XLSX source changed before final verification")
        if final_inspection.get("ready") is not True:
            raise RuntimeError("XLSX final reopen/inspection failed")
        if final_inspection.get("duplicate_row_count") != 0:
            raise RuntimeError("XLSX final verification found duplicate rows")
        if final_inspection.get("data_row_count") != len(before["retained_rows"]):
            raise RuntimeError("XLSX final verification found an unexpected row count")
        if final_inspection.get("data_fingerprint") != before["data_fingerprint"]:
            raise RuntimeError(
                "XLSX final business-value fingerprint changed after publication"
            )
        if final_inspection.get("amount_number_formats") != [number_format]:
            raise RuntimeError("XLSX final amount number_format is not uniform")
        return {
            "source_identity_before": current,
            "source_identity_after": source_after,
            "destination_identity": final_inspection["identity"],
            "source_unchanged": True,
            "destination_reopened": True,
            "sheet_name": before["sheet"].title,
            "headers": list(before["headers"]),
            "original_data_rows": len(before["rows"]),
            "removed_duplicate_rows": len(before["duplicate_indices"]),
            "retained_data_rows": len(before["retained_rows"]),
            "stable_first_row_retention": True,
            "amount_column_index": amount_col,
            "amount_number_format": number_format,
            "amount_values_semantically_unchanged": True,
            "amount_cells_numeric": True,
            "non_target_values_unchanged": True,
            "expected_data_fingerprint": final_inspection["data_fingerprint"],
        }
    except Exception:
        if temp_path.exists():
            temp_path.unlink(missing_ok=True)
        raise
